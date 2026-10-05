"""Generate plate workbooks in memory and save them only on explicit download."""
from __future__ import annotations

import hashlib
import json
import os
import re
import tempfile
import time
import uuid
from copy import deepcopy
from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

import core
from sources import LOCAL_ZONE, now_iso, read_plate

DEFAULT_CONFIGS = {
    "L2000": {
        "short_name": "L2000", "display_name": "Lipofectamine 2000",
        "reagent_label": "Lipofectamine 2000", "diluent_label": "Diluent (Opti-MEM)",
        "mode": "two_tube", "final_volume_ul": 25.0,
        "dna_to_reagent_ratio_ul_per_ug": 2.0,
        "well_overage_factor": 1.2, "bulk_overage_factor": 1.2,
    },
    "LT1": {
        "short_name": "LT1", "display_name": "LT1", "reagent_label": "LT1",
        "diluent_label": "Diluent", "mode": "single_tube", "final_volume_ul": 25.0,
        "dna_to_reagent_ratio_ul_per_ug": 3.5,
        "well_overage_factor": 1.3, "bulk_overage_factor": 1.0,
    },
}


def default_configs():
    return deepcopy(DEFAULT_CONFIGS)


def _prepare_plate(plate, tables, reagent, configs, run_warnings):
    plate_bytes = plate["bytes"]
    plate_name = Path(plate["name"]).name
    warnings = list(run_warnings)
    wells, long, parsing_warnings = core.standardize_plate_csv(read_plate(plate_bytes))
    warnings.extend(parsing_warnings)
    # Avoid accidental enormous grids caused by mistyped well coordinates.
    rows = max(core.row_sort_key(r) for r in wells["Row"])
    if rows > 64 or wells["Col"].max() > 96:
        raise core.MixMapError("The plate extends beyond 64 rows or 96 columns",
                               fixes=["Check the Well column for a mistyped plate position."])
    used = sorted(set(long["Plasmid"].map(core.clean_cell)))
    concentrations, lookup, concentration_warnings = core.load_plasmid_concentrations(
        tables, duplicate_policy="error", used_plasmids=used)
    warnings.extend(concentration_warnings)
    matched, matching_warnings = core.match_concentrations(long, lookup)
    warnings.extend(matching_warnings)
    details, summary, bulk, local_warnings = core.calculate_mix(wells, matched, configs[reagent])
    result = {"plate_name": plate_name, "plate_bytes": plate_bytes,
              "stem": re.sub(r"[^\w.\-]+", "_", Path(plate_name).stem)[:100] or "plate",
              "details": details, "summary": summary, "bulk": bulk,
              "concentrations": concentrations, "wells": len(wells),
              "entries": len(matched), "used": used,
              "warnings": core.dedupe_messages(warnings + [f"{reagent}: {w}" for w in local_warnings])}
    if plate.get("path") is not None:
        result["plate_path"] = str(plate["path"])
    return result


def _workbook_bytes(plate, reagent, configs, source_label, source_timestamp, source_hash):
    stream = BytesIO()
    core.write_mix_map_workbook(
        stream, plate["plate_name"], source_label, reagent, configs[reagent],
        plate["details"], plate["summary"], plate["bulk"], plate["concentrations"], plate["warnings"])
    # Keep provenance inside each workbook; no sidecar file is created.
    stream.seek(0)
    wb = load_workbook(stream, rich_text=True)
    try:
        ws = wb["Run config"]
        provenance = [
            ("Concentrations loaded/refreshed at", source_timestamp),
            ("Concentration snapshot SHA256", source_hash),
            ("Plate CSV SHA256", hashlib.sha256(plate["plate_bytes"]).hexdigest()),
        ]
        if "plate_path" in plate:
            provenance.append(("Plate CSV path", plate["plate_path"]))
        provenance.extend([("Generated at (Pacific time)", now_iso()),
                           ("Duplicate concentration policy", "error")])
        for label, value in provenance:
            row = ws.max_row + 1
            core.literal_cell(ws, row, 1, label)
            core.literal_cell(ws, row, 2, value)
        final = BytesIO()
        wb.save(final)
    finally:
        wb.close()
    return final.getvalue()


def generate_batch(plates, tables, source_label, source_timestamp, reagents, configs):
    """Generate one in-memory workbook per plate without an output folder.

    ``plates`` contains one to five dictionaries with ``name`` and ``bytes``;
    ``path`` is optional provenance. Validate every plate before constructing
    workbooks. Only ``save_artifact`` publishes an artifact to the filesystem.
    """
    started = time.perf_counter()
    if not isinstance(plates, (list, tuple)) or not 1 <= len(plates) <= 5:
        raise core.MixMapError(
            "Choose between 1 and 5 plate CSV files",
            fixes=["Select up to five plate layouts for each batch."],
        )
    if len(reagents) != 1 or reagents[0] not in DEFAULT_CONFIGS:
        raise core.MixMapError(
            "Choose exactly one transfectant: LT1 or L2000",
            fixes=["Select LT1 or L2000 and generate a separate workbook for each run."],
        )
    reagent = reagents[0]
    run_warnings = core.validate_run_inputs(reagents, configs, "error")
    prepared, problems, fixes = [], [], []
    for index, plate in enumerate(plates, start=1):
        if (not isinstance(plate, dict) or not isinstance(plate.get("name"), str)
                or not plate["name"].strip() or not isinstance(plate.get("bytes"), bytes)):
            raise core.MixMapError(
                "A plate CSV could not be read — no output was made",
                details=[f"Plate {index} needs a filename and CSV file contents."],
                fixes=["Choose the plate CSV files again."],
            )
        try:
            prepared.append(_prepare_plate(plate, tables, reagent, configs, run_warnings))
        except core.MixMapError as exc:
            if len(plates) == 1:
                raise
            label = f"{Path(plate['name']).name} (plate {index})"
            problems.append(f"{label}: {exc.title}")
            problems.extend(f"{label}: {detail}" for detail in exc.details)
            fixes.extend(exc.fixes)
    if problems:
        raise core.MixMapError("Cannot generate this batch — no output was made",
                               details=problems, fixes=core.dedupe_messages(fixes))

    source_hash = hashlib.sha256(json.dumps(tables, sort_keys=True).encode()).hexdigest()
    stamp = datetime.now(LOCAL_ZONE).strftime("%Y%m%d_%H%M%S")
    artifacts = []
    for plate in prepared:
        filename = f"{plate['stem']}_{reagent}_print_mixmap_{stamp}_{uuid.uuid4().hex[:8]}.xlsx"
        artifact = {"name": filename,
                    "bytes": _workbook_bytes(plate, reagent, configs, source_label,
                                             source_timestamp, source_hash),
                    "plate_name": plate["plate_name"], "reagent": reagent,
                    "summary": plate["summary"], "details": plate["details"], "bulk": plate["bulk"]}
        if "plate_path" in plate:
            artifact["plate_path"] = plate["plate_path"]
        artifacts.append(artifact)

    all_warnings = []
    for index, plate in enumerate(prepared, start=1):
        label = f"{plate['plate_name']} (plate {index}): " if len(prepared) > 1 else ""
        all_warnings.extend(label + warning for warning in plate["warnings"])
    return {"artifacts": artifacts, "plate_count": len(prepared),
            "warnings": core.dedupe_messages(all_warnings),
            "wells": sum(plate["wells"] for plate in prepared),
            "entries": sum(plate["entries"] for plate in prepared),
            "plasmids": len({name for plate in prepared for name in plate["used"]}),
            "elapsed": time.perf_counter() - started, "source": source_label,
            "source_timestamp": source_timestamp}


def generate(plate_bytes, plate_name, tables, source_label, source_timestamp,
             reagents, configs):
    """Compatibility wrapper for generating one plate workbook."""
    return generate_batch([{"name": str(plate_name), "bytes": plate_bytes}], tables, source_label,
                          source_timestamp, reagents, configs)


def save_artifact(artifact, output_dir):
    """Save exactly one generated workbook into an existing selected folder.

    The artifact is never mutated. Return the saved ``path`` and any cleanup
    ``warnings``. Publish a complete file atomically without replacing existing
    files; repeated saves receive a numbered filename. On failure, preserve the
    original exception and attach ``cleanup_warnings`` if a temporary remains.
    """
    if not output_dir or not str(output_dir).strip():
        raise core.MixMapError(
            "Choose an existing output folder",
            fixes=["Use Save Excel files to to select a folder before downloading."],
        )
    root = Path(output_dir).expanduser().resolve()
    if not root.is_dir():
        raise core.MixMapError(
            "Choose an existing output folder",
            details=[str(root)],
            fixes=["Use Save Excel files to to select a folder that already exists on your Mac."],
        )
    name = artifact.get("name")
    if (not isinstance(name, str) or not name.strip() or name in {".", ".."}
            or Path(name).name != name or "\\" in name
            or not isinstance(artifact.get("bytes"), bytes)):
        raise core.MixMapError(
            "The generated workbook is unavailable",
            fixes=["Generate the mix map again before downloading."],
        )

    temp = None
    failure = None
    cleanup_warnings = []
    try:
        descriptor, temp_name = tempfile.mkstemp(prefix=".mixmap-", suffix=".xlsx", dir=root)
        temp = Path(temp_name)
        with os.fdopen(descriptor, "wb") as handle:
            handle.write(artifact["bytes"])
            handle.flush()
            os.fsync(handle.fileno())
        for attempt in range(1000):
            candidate = name if attempt == 0 else f"{Path(name).stem}_{attempt + 1}{Path(name).suffix}"
            destination = root / candidate
            try:
                os.link(temp, destination)
            except FileExistsError:
                if attempt == 999:
                    raise
            else:
                break
    except BaseException as exc:
        failure = exc
    finally:
        if temp is not None:
            try:
                temp.unlink(missing_ok=True)
            except OSError as cleanup_error:
                cleanup_warnings.append(f"Temporary file could not be removed: {temp} ({cleanup_error}).")

    if failure is not None:
        if cleanup_warnings:
            failure.cleanup_warnings = cleanup_warnings
            failure.add_note("\n".join(cleanup_warnings))
        raise failure
    return {"path": str(destination),
            "warnings": [f"Workbook was saved. {warning}" for warning in cleanup_warnings]}
