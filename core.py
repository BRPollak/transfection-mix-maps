"""Calculation and Excel layout functions adapted from the user's Colab notebook.

The arithmetic and print layout are preserved. I/O is handled by the local app.
Non-finite numeric inputs are rejected rather than passed into calculations.
"""
import math
import os
import re
import unicodedata
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from plate_formats import get_plate_format

class MixMapError(Exception):
    def __init__(self, title, details=None, fixes=None):
        super().__init__(title)
        self.title = title
        self.details = list(details or [])
        self.fixes = list(fixes or [])


def dedupe_messages(messages):
    seen = set()
    unique = []
    for message in messages or []:
        text = clean_cell(message) if "clean_cell" in globals() else str(message).strip()
        if text and text not in seen:
            seen.add(text)
            unique.append(text)
    return unique


def format_user_message(level, title, details=None, fixes=None):
    lines = ["", f"{level}: {title}"]
    details = dedupe_messages(details)
    fixes = dedupe_messages(fixes)
    if details:
        lines.append("Details:")
        for detail in details:
            for part in str(detail).splitlines() or [""]:
                lines.append(f"- {part}")
    if fixes:
        lines.append("How to fix:")
        for fix in fixes:
            for part in str(fix).splitlines() or [""]:
                lines.append(f"- {part}")
    return "\n".join(lines)


def print_user_error(error):
    print(format_user_message("ERROR", error.title, error.details, error.fixes))


def print_user_warnings(warnings):
    warnings = dedupe_messages(warnings)
    if warnings:
        print(
            format_user_message(
                "WARNING",
                "Review these non-fatal issues before pipetting",
                warnings,
                ["Check the Warnings and Run config sheets in the output workbook."],
            )
        )


def summarize_list(values, max_items=8):
    values = [clean_cell(v) for v in values if clean_cell(v)] if "clean_cell" in globals() else [str(v) for v in values]
    values = dedupe_messages(values)
    if len(values) <= max_items:
        return ", ".join(values)
    return ", ".join(values[:max_items]) + f", and {len(values) - max_items} more"


def limited_examples(items, max_items=12):
    items = [str(item) for item in items if str(item).strip()]
    if len(items) <= max_items:
        return items
    return items[:max_items] + [f"... {len(items) - max_items} more not shown"]


def plasmid_is_csv_relevant(plasmid, used_plasmids=None):
    if not used_plasmids:
        return True
    name = clean_cell(plasmid)
    if not name:
        return False
    used_exact = {clean_cell(item) for item in used_plasmids if clean_cell(item)}
    used_lower = {item.lower() for item in used_exact}
    return name in used_exact or name.lower() in used_lower


def validate_run_inputs(REAGENTS_TO_RUN, REAGENT_CONFIGS, DUPLICATE_CONCENTRATION_POLICY):
    warnings = []
    if not REAGENTS_TO_RUN:
        raise MixMapError(
            "No reagent configuration selected",
            details=["reagents_to_run did not contain L2000 or LT1."],
            fixes=["Choose L2000 or LT1 from the transfectant selector."],
        )
    unknown_reagents = [name for name in REAGENTS_TO_RUN if name not in REAGENT_CONFIGS]
    if unknown_reagents:
        raise MixMapError(
            "Unknown reagent configuration",
            details=[f"Unknown selections: {summarize_list(unknown_reagents)}"],
            fixes=[f"Use one of: {', '.join(sorted(REAGENT_CONFIGS))}."],
        )
    if DUPLICATE_CONCENTRATION_POLICY not in {"error", "first", "last"}:
        raise MixMapError(
            "Invalid duplicate concentration policy",
            details=[f"Got: {DUPLICATE_CONCENTRATION_POLICY!r}"],
            fixes=["Use error, first, or last."],
        )

    numeric_errors = []
    for config_name in REAGENTS_TO_RUN:
        config = REAGENT_CONFIGS[config_name]
        checks = [
            ("final_volume_ul", config["final_volume_ul"], "must be greater than 0"),
            ("dna_to_reagent_ratio_ul_per_ug", config["dna_to_reagent_ratio_ul_per_ug"], "must be greater than 0"),
            ("well_overage_factor", config["well_overage_factor"], "must be greater than 0"),
            ("bulk_overage_factor", config.get("bulk_overage_factor", 1.0), "must be greater than 0"),
        ]
        for setting, value, requirement in checks:
            if not isinstance(value, (int, float)) or not math.isfinite(float(value)) or float(value) <= 0:
                numeric_errors.append(f"{config_name} {setting}={value!r} {requirement}.")
        if isinstance(config.get("well_overage_factor"), (int, float)) and 0 < config["well_overage_factor"] < 1:
            warnings.append(
                f"{config_name} well_overage_factor is below 1.0; this creates less mix than the nominal per-well volume."
            )
        if isinstance(config.get("bulk_overage_factor"), (int, float)) and 0 < config["bulk_overage_factor"] < 1:
            warnings.append(
                f"{config_name} bulk_overage_factor is below 1.0; the bulk reagent tube may be short."
            )
    if numeric_errors:
        raise MixMapError(
            "Invalid reagent volume settings",
            details=numeric_errors,
            fixes=["Correct the Reagent settings in the app and rerun."],
        )
    return warnings


def normalize_header(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = text.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def clean_cell(value):
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return ""
    return str(value).strip()


def parse_float(value):
    """Read a complete DNA mass in ng, never extract digits from other text."""
    match = re.fullmatch(
        r"([+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?)\s*(?:ng)?",
        clean_cell(value), flags=re.IGNORECASE,
    )
    if match is None:
        return None
    number = float(match.group(1))
    return number if math.isfinite(number) else None


def extract_spreadsheet_id(text):
    if not text:
        return None
    text = str(text).strip()
    patterns = [
        r"/spreadsheets/d/([A-Za-z0-9_-]+)",
        r"(?:spreadsheet:)([A-Za-z0-9_-]+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return match.group(1)
    if re.fullmatch(r"[A-Za-z0-9_-]{25,}", text):
        return text
    return None


def values_to_dataframe(values):
    if not values:
        return pd.DataFrame()

    header_idx = None
    # Prefer a recognizable stock header to a title or notes row. A sheet with
    # only a Plasmid column still tells us that a name exists without a stock
    # concentration, which is different from an absent plasmid.
    for idx, row in enumerate(values):
        name_col = find_plasmid_name_column(row)
        conc_col = find_concentration_column(row)
        if name_col is not None and conc_col is not None and name_col != conc_col:
            header_idx = idx
            break
    if header_idx is None:
        for idx, row in enumerate(values):
            name_col = find_plasmid_name_column(row)
            if name_col is not None and normalize_header(name_col) in {"plasmid", "plasmidname", "name"}:
                header_idx = idx
                break
    if header_idx is None:
        for idx, row in enumerate(values):
            nonblank = [clean_cell(v) for v in row if clean_cell(v)]
            if len(nonblank) >= 2:
                header_idx = idx
                break
    if header_idx is None:
        return pd.DataFrame()

    headers = [clean_cell(h) or f"Column {i + 1}" for i, h in enumerate(values[header_idx])]
    rows = values[header_idx + 1 :]
    width = len(headers)
    padded_rows = [(list(row) + [""] * width)[:width] for row in rows]
    df = pd.DataFrame(padded_rows, columns=headers)
    df.attrs["header_row"] = header_idx + 1
    return df


def find_plasmid_name_column(columns):
    normalized = {col: normalize_header(col) for col in columns}
    for col, norm in normalized.items():
        if norm in {"plasmidname", "plasmid", "name"}:
            return col
    for col, norm in normalized.items():
        if "plasmid" in norm and "name" in norm:
            return col
    for col, norm in normalized.items():
        if "plasmid" in norm:
            return col
    return None


def normalize_unit_header(value):
    return normalize_header(str(value).replace("µ", "u").replace("μ", "u"))


def concentration_columns(columns):
    """Keep every candidate, including duplicate labels, for ambiguity checks."""
    candidates = []
    for col in columns:
        norm = normalize_unit_header(col)
        if ("concentration" in norm or re.search(r"\bconc\b", str(col), re.I)
                or norm.startswith("concng")
                or norm in {"ngul", "ngperul", "ngmicroliter", "ngpermicroliter"}):
            candidates.append(col)
    return candidates


def find_concentration_column(columns):
    candidates = concentration_columns(columns)
    return candidates[0] if candidates else None


def concentration_header_supported(column):
    norm = normalize_unit_header(column)
    unit = re.sub(r"(?:concentrations?|conc|plasmid|dna|stock)", "", norm)
    return unit in {"", "ngul", "ngperul", "ngmicroliter", "ngpermicroliter"}


def parse_concentration(value):
    """Read finite positive stock concentrations, optionally labeled ng/uL."""
    match = re.fullmatch(
        r"([+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?)\s*(?:ng\s*(?:/|per)\s*[uµμ]l)?",
        str(value).strip(), flags=re.IGNORECASE,
    )
    if match is None:
        return None
    result = float(match.group(1))
    return result if math.isfinite(result) and result > 0 else None


def scan_concentration_tables(tables):
    """Collect usable stocks and named rows without rejecting an entire Sheet.

    Entries retain unusable named rows so plate validation can distinguish an
    absent name from an existing name whose concentration cannot be used.
    """
    records, entries, worksheets = [], {}, []
    for title, values in (tables.items() if isinstance(tables, dict) else []):
        item = {"title": title, "status": "skipped", "header_row": None,
                "plasmid_column": None, "concentration_column": None,
                "valid_row_count": 0, "errors": [], "warnings": []}
        worksheets.append(item)
        try:
            frame = values_to_dataframe(values)
        except (TypeError, ValueError):
            continue
        columns = list(frame.columns)
        name_col = find_plasmid_name_column(columns)
        conc_candidates = concentration_columns(columns)
        conc_col = conc_candidates[0] if conc_candidates else None
        item.update(header_row=int(frame.attrs.get("header_row", 1)),
                    plasmid_column=name_col, concentration_column=conc_col)
        if name_col is None:
            continue
        mapping_issue = None
        if conc_col is None:
            mapping_issue = "concentration column is missing or unrecognized"
        elif name_col == conc_col or len(conc_candidates) != 1:
            mapping_issue = "concentration column is ambiguous"
        elif not concentration_header_supported(conc_col):
            mapping_issue = "concentration units must be ng/µL"
        name_positions = [index for index, column in enumerate(columns) if column == name_col]
        if len(name_positions) != 1:
            mapping_issue = "plasmid column is ambiguous"
        for row_index, row in frame.iterrows():
            source_row = int(row_index) + item["header_row"] + 1
            names = {clean_cell(row.iloc[index]) for index in name_positions}
            for name in sorted(names - {""}):
                raw = "" if mapping_issue else clean_cell(row.get(conc_col))
                concentration = None if mapping_issue else parse_concentration(raw)
                issue = mapping_issue
                if issue is None and concentration is None:
                    issue = ("concentration is missing" if not raw else
                             "concentration must be a finite, positive number in ng/µL")
                entry = {"Plasmid": name, "Concentration_ng_per_uL": concentration,
                         "Source worksheet": title, "Source row": source_row, "issue": issue}
                entries.setdefault(name, []).append(entry)
                if issue is None:
                    record = {key: value for key, value in entry.items() if key != "issue"}
                    records.append(record)
                    item["valid_row_count"] += 1
        if item["valid_row_count"]:
            item["status"] = "available"
    return {"records": records, "entries": entries, "worksheets": worksheets}


def load_plasmid_concentrations(tables, duplicate_policy="error", used_plasmids=None):
    scan = scan_concentration_tables(tables)
    entries = scan["entries"]
    requested = sorted(entries if used_plasmids is None else
                       {clean_cell(item) for item in used_plasmids if clean_cell(item)})
    lookup, warnings, problems = {}, [], []
    selected_records = []
    folded = {}
    for name in entries:
        folded.setdefault(name.lower(), []).append(name)
    for plasmid in requested:
        candidates = [plasmid] if plasmid in entries else folded.get(plasmid.lower(), [])
        if not candidates:
            problems.append(f"{plasmid}: not found in the Google Sheet.")
            continue
        if len(candidates) > 1:
            problems.append(f"{plasmid}: ambiguous name match ({summarize_list(candidates, 4)}); use the exact Sheet name.")
            continue
        name = candidates[0]
        valid = [entry for entry in entries[name] if entry["issue"] is None]
        if not valid:
            reasons = "; ".join(dedupe_messages(entry["issue"] for entry in entries[name]))
            problems.append(f"{plasmid}: {reasons}.")
            continue
        concentrations = {entry["Concentration_ng_per_uL"] for entry in valid}
        conflict = max(concentrations) - min(concentrations) > 1e-9
        if conflict and duplicate_policy == "error":
            values = summarize_list([f"{value:g}" for value in sorted(concentrations)], 4)
            problems.append(f"{plasmid}: conflicting concentrations ({values} ng/µL); one concentration is required.")
            continue
        if conflict:
            warnings.append(f"Resolved conflicting concentrations for '{plasmid}' using policy '{duplicate_policy}'.")
        chosen = valid[-1] if duplicate_policy == "last" else valid[0]
        lookup[name] = {key: value for key, value in chosen.items() if key != "issue"}
        selected_records.extend({key: value for key, value in entry.items() if key != "issue"}
                                for entry in valid)
    if problems:
        raise MixMapError(
            "Cannot generate mix maps — no output was made",
            details=problems,
            fixes=["Add or correct these plasmid concentrations in the Google Sheet, refresh it, and generate again."],
        )
    if not lookup:
        raise MixMapError(
            "Cannot generate mix maps — no output was made",
            details=["No plasmid concentrations are available for this plate."],
            fixes=["Add the plate's plasmids and positive concentrations in ng/µL to the Google Sheet."],
        )
    concentration_df = pd.DataFrame(selected_records).drop_duplicates().sort_values(
        ["Plasmid", "Source worksheet", "Source row"])
    return concentration_df, lookup, warnings


def find_well_column(columns):
    for col in columns:
        if normalize_header(col) == "well":
            return col
    for col in columns:
        if "well" in normalize_header(col):
            return col
    return None


def is_mass_column(col):
    norm = normalize_header(col)
    return (
        "mass" in norm
        or "totalng" in norm
        or "desiredng" in norm
        or norm in {"ng", "dnang", "amountng", "nanograms"}
    ) and "plasmid" not in norm


def slot_number(col):
    match = re.search(r"(\d+)", normalize_header(col))
    return match.group(1) if match else None


def validate_mass_column(column):
    norm = normalize_unit_header(column)
    units = re.sub(r"(?:mass|total|desired|amount|dna|\d+)", "", norm)
    if units not in {"", "ng", "nanogram", "nanograms"}:
        raise MixMapError(
            "Unsupported DNA mass units in plate CSV",
            details=[f"Mass column {column!r} must contain masses in ng."],
            fixes=["Convert DNA masses to ng and label the column Mass (ng), or Mass1 (ng) for wide format."],
        )


def parse_well(well):
    text = clean_cell(well).upper().replace(" ", "")
    match = re.fullmatch(r"([A-Z]+)0*([0-9]+)", text)
    if not match or int(match.group(2)) < 1:
        raise ValueError(f"Malformed well value: {well!r}")
    return match.group(1), int(match.group(2))


def parse_well_or_raise(well, csv_row=None, column_name="Well"):
    try:
        return parse_well(well)
    except ValueError as exc:
        row_text = f"CSV row {csv_row}" if csv_row is not None else "CSV input"
        raise MixMapError(
            "Malformed well value in plate CSV",
            details=[f"{row_text}, column '{column_name}': {well!r}"],
            fixes=["Use well IDs like A1, A01, B12, H8, etc. Do not include ranges or plate names in the Well column."],
        ) from exc


def row_sort_key(label):
    total = 0
    for ch in str(label):
        if "A" <= ch <= "Z":
            total = total * 26 + (ord(ch) - ord("A") + 1)
    return total


BULK_MIX_SYMBOLS = ("▲", "●", "■", "◆", "★")


def group_dna_masses(long_df):
    """Rank positive DNA totals consistently for the preview and bulk recipes.

    Decimal sums make 0.1 + 0.2 and 0.3 one group without rounding distinct
    requested totals together. Numbering restarts for each complete plate.
    """
    if long_df is None or long_df.empty:
        return []
    totals = {}
    for record in long_df.to_dict("records"):
        row, col = parse_well(record["Well"])
        try:
            mass = Decimal(str(record["Mass_ng"]))
        except (InvalidOperation, ValueError):
            continue
        if not mass.is_finite() or mass <= 0:
            continue
        position = (row, col)
        totals[position] = totals.get(position, Decimal(0)) + mass

    grouped = {}
    for position, mass in totals.items():
        grouped.setdefault(mass, []).append(position)
    physical_position = lambda position: (row_sort_key(position[0]), position[1])
    ordered = sorted(grouped.items(), key=lambda item: (-len(item[1]), min(map(physical_position, item[1]))))
    groups = []
    for number, (mass, positions) in enumerate(ordered, 1):
        symbol = ("" if len(ordered) == 1 else BULK_MIX_SYMBOLS[number - 1]
                  if len(ordered) <= len(BULK_MIX_SYMBOLS) else str(number))
        groups.append({
            "Number": number,
            "Total DNA_ng": mass,
            "Wells": [f"{row}{col}" for row, col in sorted(positions, key=physical_position)],
            "Symbol": symbol,
        })
    return groups


def row_label_from_index(index):
    letters = []
    while index > 0:
        index, remainder = divmod(index - 1, 26)
        letters.append(chr(ord("A") + remainder))
    return "".join(reversed(letters))


def contiguous_rows(observed_rows):
    observed_rows = sorted(set(observed_rows), key=row_sort_key)
    if not observed_rows:
        return []
    if all(re.fullmatch(r"[A-Z]+", row) for row in observed_rows):
        start = row_sort_key(observed_rows[0])
        end = row_sort_key(observed_rows[-1])
        if end - start <= 384:
            return [row_label_from_index(i) for i in range(start, end + 1)]
    return observed_rows


def detect_wide_pairs(columns):
    plasmid_cols = {}
    mass_cols = {}

    for col in columns:
        norm = normalize_header(col)
        match = re.fullmatch(r"plasmid(\d+)", norm)
        if match:
            plasmid_cols[match.group(1)] = col

    for col in columns:
        if is_mass_column(col):
            slot = slot_number(col)
            if slot:
                mass_cols.setdefault(slot, []).append(col)

    if not plasmid_cols:
        return []

    pairs = []
    unnumbered_mass_cols = [col for col in columns if is_mass_column(col) and slot_number(col) is None]
    for slot, plasmid_col in sorted(plasmid_cols.items(), key=lambda item: int(item[0])):
        candidates = list(mass_cols.get(slot, []))
        if len(plasmid_cols) == 1:
            candidates.extend(unnumbered_mass_cols)
        if len(candidates) > 1:
            raise MixMapError(
                "Ambiguous DNA mass columns in wide-format CSV",
                details=[f"Slot {slot} has multiple mass columns: {', '.join(repr(col) for col in candidates)}."],
                fixes=[f"Keep exactly one Mass{slot} (ng) column for Plasmid{slot}."],
            )
        if not candidates:
            raise MixMapError(
                "Wide-format CSV has a plasmid column without a matching mass column",
                details=[f"Found {plasmid_col!r} but no matching mass column for slot {slot}."],
                fixes=[f"Add a Mass{slot} (ng) column, or rename the existing mass column so it has the same slot number."],
            )
        pairs.append((slot, plasmid_col, candidates[0]))
    return pairs


def detect_long_columns(columns):
    well_col = find_well_column(columns)
    plasmid_col = None

    for col in columns:
        norm = normalize_header(col)
        if norm in {"plasmid", "plasmidname", "construct", "dna"} or ("plasmid" in norm and "name" in norm):
            plasmid_col = col
            break
    mass_columns = [col for col in columns if is_mass_column(col)]
    if well_col and plasmid_col and mass_columns:
        if len(mass_columns) > 1:
            raise MixMapError(
                "Ambiguous DNA mass columns in long-format CSV",
                details=[f"Found multiple mass columns: {', '.join(repr(col) for col in mass_columns)}."],
                fixes=["Keep exactly one Mass (ng) column for the requested plasmid mass."],
            )
        return well_col, plasmid_col, mass_columns[0]
    return None


def standardize_plate_csv(raw_df, *, plate_type=None):
    """Read DNA entries, optionally enforcing the selected physical plate layout.

    ``None`` retains unrestricted coordinates for internal calculation callers.
    Empty and valid-zero template wells outside an explicit format are ignored.
    """
    plate_format = get_plate_format(plate_type) if plate_type is not None else None
    warnings = []
    if raw_df.empty:
        raise MixMapError(
            "Plate CSV has no data rows",
            details=["The file loaded, but pandas found no rows."],
            fixes=["Use a CSV with one row per well or one row per plasmid/well entry."],
        )

    well_col = find_well_column(raw_df.columns)
    if well_col is None:
        raise MixMapError(
            "CSV must include a Well column",
            details=[f"Columns found: {summarize_list(raw_df.columns, 20)}"],
            fixes=["Add a Well column with values like A1, A2, B1, etc."],
        )

    well_records = []
    seen_wells = set()
    well_occurrences = {}
    canonical_wells = []
    supported_wells = set()
    for csv_row_index, value in raw_df[well_col].items():
        well = clean_cell(value)
        if not well:
            canonical_wells.append("")
            continue
        row_label, col_index = parse_well_or_raise(well, csv_row=int(csv_row_index) + 2, column_name=well_col)
        well = f"{row_label}{col_index}"
        canonical_wells.append(well)
        seen_wells.add(well)
        if plate_format is None or plate_format.contains(row_label, col_index):
            well_occurrences.setdefault(well, []).append(int(csv_row_index) + 2)
            if well not in supported_wells:
                supported_wells.add(well)
                well_records.append({"Well": well, "Row": row_label, "Col": col_index})
    raw_df = raw_df.copy()
    raw_df[well_col] = canonical_wells

    if not seen_wells:
        raise MixMapError(
            "No valid wells were found in the CSV",
            details=[f"Well column: {well_col}"],
            fixes=["Fill the Well column with values like A1, A2, B1, etc."],
        )

    wide_pairs = detect_wide_pairs(raw_df.columns)
    long_records = []
    outside_entries = []

    if wide_pairs:
        for _, _, mass_col in wide_pairs:
            validate_mass_column(mass_col)
        duplicate_wells = {well: rows for well, rows in well_occurrences.items() if len(rows) > 1}
        if duplicate_wells:
            details = [f"{well}: CSV rows {', '.join(map(str, rows))}" for well, rows in list(duplicate_wells.items())[:20]]
            if len(duplicate_wells) > 20:
                details.append(f"... {len(duplicate_wells) - 20} more duplicate wells not shown")
            raise MixMapError(
                "Duplicate well rows found in wide-format CSV",
                details=details,
                fixes=[
                    "Keep one row per well when using Plasmid1/Mass1, Plasmid2/Mass2 wide columns.",
                    "If you intentionally need multiple rows per well, use long format columns: Well, Plasmid, Mass (ng).",
                ],
            )
        for csv_row_index, row in raw_df.iterrows():
            well = row[well_col]
            if not well:
                if any(clean_cell(value) for value in row):
                    raise MixMapError(
                        "Missing well value in wide-format CSV",
                        details=[f"CSV row {int(csv_row_index) + 2} contains data but has no well."],
                        fixes=["Fill the Well column for every populated row, or remove the entire row."],
                    )
                continue
            for slot, plasmid_col, mass_col in wide_pairs:
                plasmid = clean_cell(row.get(plasmid_col))
                raw_mass = clean_cell(row.get(mass_col))
                mass_ng = parse_float(row.get(mass_col))
                if not plasmid and not raw_mass:
                    continue
                if plasmid and mass_ng is None:
                    raise MixMapError(
                        "Missing or unreadable DNA mass in plate CSV",
                        details=[f"CSV row {int(csv_row_index) + 2}, well {well}, plasmid {plasmid!r}, mass column {mass_col!r}: {raw_mass!r}"],
                        fixes=["Enter a numeric mass in ng for every plasmid name."],
                    )
                if not plasmid and raw_mass:
                    raise MixMapError(
                        "Mass is present without a plasmid name in plate CSV",
                        details=[f"CSV row {int(csv_row_index) + 2}, well {well}, slot {slot}, mass column {mass_col!r}: {raw_mass!r}"],
                        fixes=["Add the plasmid name, or clear the mass cell if this slot should be empty."],
                    )
                if mass_ng < 0:
                    raise MixMapError(
                        "Negative DNA mass in plate CSV",
                        details=[f"CSV row {int(csv_row_index) + 2}, well {well}, plasmid {plasmid!r}: {mass_ng} ng"],
                        fixes=["Mass values must be zero or positive; correct the CSV and rerun."],
                    )
                if mass_ng == 0:
                    if well in supported_wells:
                        warnings.append(f"CSV row {int(csv_row_index) + 2}, well {well}, plasmid {plasmid!r} has 0 ng and was skipped.")
                    continue
                if well not in supported_wells:
                    outside_entries.append((int(csv_row_index) + 2, well))
                    continue
                long_records.append(
                    {
                        "Well": well,
                        "Slot": int(slot),
                        "Plasmid": plasmid,
                        "Mass_ng": float(mass_ng),
                        "CSV row": int(csv_row_index) + 2,
                    }
                )
    else:
        long_cols = detect_long_columns(raw_df.columns)
        if long_cols is None:
            raise MixMapError(
                "Could not detect plasmid and mass columns in CSV",
                details=[f"Columns found: {summarize_list(raw_df.columns, 20)}"],
                fixes=[
                    "Use wide columns like Well, Plasmid1, Mass1 (ng), Plasmid2, Mass2 (ng).",
                    "Or use long columns like Well, Plasmid, Mass (ng).",
                ],
            )
        well_col, plasmid_col, mass_col = long_cols
        validate_mass_column(mass_col)
        slot_counter = {}
        for csv_row_index, row in raw_df.iterrows():
            row_num = int(csv_row_index) + 2
            well = row[well_col]
            plasmid = clean_cell(row.get(plasmid_col))
            raw_mass = clean_cell(row.get(mass_col))
            mass_ng = parse_float(row.get(mass_col))
            if not well and not plasmid and not raw_mass:
                continue
            if not well:
                raise MixMapError(
                    "Missing well value in long-format CSV",
                    details=[f"CSV row {row_num}: plasmid={plasmid!r}, mass={raw_mass!r}"],
                    fixes=["Fill the Well column for every plasmid/mass row."],
                )
            if not plasmid and not raw_mass:
                continue
            if not plasmid:
                raise MixMapError(
                    "Missing plasmid name in long-format CSV",
                    details=[f"CSV row {row_num}, well {well}: mass={raw_mass!r}"],
                    fixes=["Fill the plasmid column, or remove the row if it should not be included."],
                )
            if mass_ng is None:
                raise MixMapError(
                    "Missing or unreadable DNA mass in long-format CSV",
                    details=[f"CSV row {row_num}, well {well}, plasmid {plasmid!r}, mass={raw_mass!r}"],
                    fixes=["Enter a numeric mass in ng for every plasmid row."],
                )
            if mass_ng < 0:
                raise MixMapError(
                    "Negative DNA mass in long-format CSV",
                    details=[f"CSV row {row_num}, well {well}, plasmid {plasmid!r}: {mass_ng} ng"],
                    fixes=["Mass values must be zero or positive; correct the CSV and rerun."],
                )
            if mass_ng == 0:
                if well in supported_wells:
                    warnings.append(f"CSV row {row_num}, well {well}, plasmid {plasmid!r} has 0 ng and was skipped.")
                continue
            if well not in supported_wells:
                outside_entries.append((row_num, well))
                continue
            slot_counter[well] = slot_counter.get(well, 0) + 1
            long_records.append(
                {
                    "Well": well,
                    "Slot": slot_counter[well],
                    "Plasmid": plasmid,
                    "Mass_ng": float(mass_ng),
                    "CSV row": row_num,
                }
            )

    if outside_entries:
        raise MixMapError(
            "DNA is assigned outside the selected plate type",
            details=dedupe_messages([
                f"CSV row {row_num}, well {well} contains positive DNA mass; "
                f"the selected {plate_format.label} plate allows {plate_format.well_range}."
                for row_num, well in outside_entries
            ]),
            fixes=["Choose the correct plate type, or move/remove DNA assigned to unsupported wells."],
        )
    wells_df = pd.DataFrame(well_records, columns=["Well", "Row", "Col"]).drop_duplicates("Well")
    long_df = pd.DataFrame(long_records)
    if long_df.empty:
        raise MixMapError(
            "No usable plasmid/mass entries were found in the CSV",
            details=warnings or ["The CSV contains wells, but all plasmid/mass entries were blank or zero."],
            fixes=["Add at least one plasmid name with a positive mass in ng and rerun."],
        )
    return wells_df, long_df, warnings


def match_concentrations(long_df, concentration_lookup):
    if long_df.empty:
        raise MixMapError(
            "No plasmid rows were available to match concentrations",
            fixes=["Check the plate CSV parsing warnings and rerun."],
        )

    lower_lookup = {}
    for name in concentration_lookup:
        lower_lookup.setdefault(name.lower(), []).append(name)

    records = []
    warnings = []
    missing = []
    ambiguous = []
    for _, row in long_df.iterrows():
        plasmid = row["Plasmid"]
        matched_name = plasmid
        if plasmid not in concentration_lookup:
            candidates = lower_lookup.get(str(plasmid).lower(), [])
            if len(candidates) == 1:
                matched_name = candidates[0]
                warnings.append(
                    f"CSV plasmid '{plasmid}' in well {row['Well']} matched concentration table name '{matched_name}' by case-insensitive lookup."
                )
            elif len(candidates) > 1:
                ambiguous.append((row, candidates))
                continue
            else:
                missing.append(row)
                continue

        conc_record = concentration_lookup[matched_name]
        out = row.to_dict()
        out["Matched plasmid"] = matched_name
        out["Concentration_ng_per_uL"] = conc_record["Concentration_ng_per_uL"]
        out["Concentration source worksheet"] = conc_record["Source worksheet"]
        out["Concentration source row"] = conc_record["Source row"]
        records.append(out)

    if missing or ambiguous:
        details = []
        if missing:
            missing_df = pd.DataFrame(missing)
            for plasmid, group in missing_df.groupby("Plasmid", sort=True):
                wells = summarize_list(group["Well"].tolist(), 12)
                csv_rows = summarize_list([str(v) for v in group["CSV row"].tolist()], 12)
                details.append(
                    f"Missing concentration for '{plasmid}': referenced in wells {wells}; CSV rows {csv_rows}."
                )
        for row, candidates in ambiguous[:20]:
            details.append(
                f"Ambiguous case-insensitive match for '{row['Plasmid']}' in well {row['Well']} (CSV row {row['CSV row']}): candidates {summarize_list(candidates)}."
            )
        if len(ambiguous) > 20:
            details.append(f"... {len(ambiguous) - 20} more ambiguous plasmid matches not shown")
        raise MixMapError(
            "Missing or ambiguous plasmid concentrations",
            details=details,
            fixes=[
                "Add a positive numeric concentration for each missing plasmid in the concentration table.",
                "Make plasmid names match exactly between the CSV and concentration table, including suffixes and spacing.",
                "Review any skipped concentration table row warnings; a concentration row may have been ignored because the name or concentration was blank/unreadable.",
            ],
        )

    if not records:
        raise MixMapError(
            "No plasmids could be matched to concentrations",
            fixes=["Check plasmid names and concentration rows in the concentration table."],
        )
    return pd.DataFrame(records), warnings


def describe_well_components(well_name, details, max_items=3):
    subset = details[details["Well"] == well_name].sort_values("Working DNA volume_uL", ascending=False)
    parts = []
    for _, item in subset.head(max_items).iterrows():
        parts.append(
            f"{item['Matched plasmid']} {item['Working DNA volume_uL']:.3f} uL "
            f"({item['Mass_ng']:.1f} ng / {item['Concentration_ng_per_uL']:.3f} ng/uL x overage)"
        )
    return "; ".join(parts)


def _tiny_dna_warning(details, volume_formatter=None):
    """Render the same diagnostic from original values for tables or print."""
    if details.empty:
        return None
    tiny_dna = details[(details["Working DNA volume_uL"] > 0) & (details["Working DNA volume_uL"] < 0.2)]
    if tiny_dna.empty:
        return None
    if volume_formatter is None:
        volume_formatter = lambda value: f"{value:.3f} uL"
    examples = [
        f"{row['Well']} {row['Matched plasmid']} {volume_formatter(row['Working DNA volume_uL'])}"
        for _, row in tiny_dna.head(12).iterrows()
    ]
    return (
        f"{len(tiny_dna)} DNA component volume(s) are below 0.200 uL and may be difficult to pipette accurately. "
        f"Examples: {summarize_list(examples, 12)}."
    )


def calculate_mix(wells_df, matched_df, config):
    warnings = []
    details = matched_df.copy()
    if details.empty:
        raise MixMapError(
            "No matched plasmid entries are available for calculation",
            fixes=["Check the CSV and concentration table concentration matching steps."],
        )
    if (details["Concentration_ng_per_uL"] <= 0).any():
        bad = details[details["Concentration_ng_per_uL"] <= 0]
        raise MixMapError(
            "Invalid plasmid concentration detected after matching",
            details=[f"{row['Matched plasmid']} in well {row['Well']}: {row['Concentration_ng_per_uL']} ng/uL" for _, row in bad.head(20).iterrows()],
            fixes=["Use positive numeric concentrations in the concentration table."],
        )

    details["Base DNA volume_uL"] = details["Mass_ng"] / details["Concentration_ng_per_uL"]
    details["Working DNA volume_uL"] = details["Base DNA volume_uL"] * config["well_overage_factor"]

    mass_groups = group_dna_masses(details)
    mass_keys = {well: group["Total DNA_ng"] for group in mass_groups for well in group["Wells"]}
    summary_rows = []
    for _, well in wells_df.iterrows():
        well_name = well["Well"]
        subset = details[details["Well"] == well_name] if not details.empty else pd.DataFrame()
        total_mass = float(mass_keys.get(well_name, Decimal(0)))
        total_dna_working = float(subset["Working DNA volume_uL"].sum()) if not subset.empty else 0.0
        reagent_working = (
            total_mass / 1000.0
            * config["dna_to_reagent_ratio_ul_per_ug"]
            * config["well_overage_factor"]
        )

        mode = config.get("mode", "two_tube")
        if mode == "two_tube":
            dna_mix_target = config["final_volume_ul"] * config["well_overage_factor"] / 2.0
            transfection_mix_target = dna_mix_target
            dna_diluent = dna_mix_target - total_dna_working
            transfection_diluent = transfection_mix_target - reagent_working
        elif mode == "single_tube":
            dna_mix_target = config["final_volume_ul"] * config["well_overage_factor"]
            transfection_mix_target = 0.0
            dna_diluent = dna_mix_target - total_dna_working - reagent_working
            transfection_diluent = 0.0
        else:
            raise MixMapError(
                "Unknown reagent mode",
                details=[f"Mode for {config.get('display_name')}: {mode!r}"],
                fixes=["Use reagent mode two_tube or single_tube in REAGENT_CONFIGS."],
            )

        if total_mass == 0:
            dna_mix_target = transfection_mix_target = 0.0
            dna_diluent = transfection_diluent = 0.0

        summary_rows.append(
            {
                "Well": well_name,
                "Row": well["Row"],
                "Col": int(well["Col"]),
                "Total DNA_ng": total_mass,
                "Total working DNA_uL": total_dna_working,
                "DNA mix target_uL": dna_mix_target,
                "DNA diluent_uL": dna_diluent,
                "Reagent_uL": reagent_working,
                "Transfection mix target_uL": transfection_mix_target,
                "Transfection diluent_uL": transfection_diluent,
            }
        )

    summary = pd.DataFrame(summary_rows)

    numeric_columns = [column for column in summary if column.endswith(("_ng", "_uL"))]
    if (not all(math.isfinite(value) for value in details["Working DNA volume_uL"])
            or not all(math.isfinite(value) for value in summary[numeric_columns].to_numpy().flat)):
        raise MixMapError(
            "The inputs produce non-finite calculated volumes",
            fixes=["Check DNA masses, stock concentrations, and reagent settings for extreme values."],
        )

    bad_dna = summary[(summary["Total DNA_ng"] > 0) & (summary["DNA diluent_uL"] < -1e-9)]
    bad_transfection = summary[
        (summary["Total DNA_ng"] > 0) & (summary["Transfection diluent_uL"] < -1e-9)
    ]
    if not bad_dna.empty or not bad_transfection.empty:
        details_lines = []
        for _, row in bad_dna.head(20).iterrows():
            if config.get("mode") == "single_tube":
                details_lines.append(
                    f"Well {row['Well']}: DNA + reagent needs {row['Total working DNA_uL'] + row['Reagent_uL']:.3f} uL, "
                    f"but complete mix target is {row['DNA mix target_uL']:.3f} uL; over by {-row['DNA diluent_uL']:.3f} uL. "
                    f"Largest DNA contributors: {describe_well_components(row['Well'], details)}"
                )
            else:
                details_lines.append(
                    f"Well {row['Well']}: DNA components need {row['Total working DNA_uL']:.3f} uL, "
                    f"but DNA mix target is {row['DNA mix target_uL']:.3f} uL; over by {-row['DNA diluent_uL']:.3f} uL. "
                    f"Largest DNA contributors: {describe_well_components(row['Well'], details)}"
                )
        for _, row in bad_transfection.head(20).iterrows():
            details_lines.append(
                f"Well {row['Well']}: {config['reagent_label']} needs {row['Reagent_uL']:.3f} uL, "
                f"but transfection mix target is {row['Transfection mix target_uL']:.3f} uL; over by {-row['Transfection diluent_uL']:.3f} uL."
            )
        hidden = max(len(bad_dna) - 20, 0) + max(len(bad_transfection) - 20, 0)
        if hidden:
            details_lines.append(f"... {hidden} more volume-limit issue(s) not shown")
        raise MixMapError(
            "Some wells need more volume than the configured mix allows",
            details=details_lines,
            fixes=[
                "Increase final_volume_ul.",
                "Reduce requested DNA mass for the listed wells.",
                "Use a more concentrated plasmid prep for the largest DNA contributors.",
                "For L2000/two-tube mode, remember only half of final_volume_ul is available for the DNA/diluent tube before overage.",
            ],
        )

    low_dna_diluent = summary[
        (summary["Total DNA_ng"] > 0)
        & (summary["DNA diluent_uL"] >= 0)
        & (summary["DNA diluent_uL"] < 1.0)
    ]
    if not low_dna_diluent.empty:
        warnings.append(
            f"{len(low_dna_diluent)} well(s) leave less than 1.000 uL DNA diluent, which may be hard to pipette accurately. Examples: {summarize_list(low_dna_diluent['Well'].tolist(), 12)}."
        )
    low_transfection_diluent = summary[
        (summary["Transfection mix target_uL"] > 0)
        & (summary["Transfection diluent_uL"] >= 0)
        & (summary["Transfection diluent_uL"] < 1.0)
    ]
    if not low_transfection_diluent.empty:
        warnings.append(
            f"{len(low_transfection_diluent)} well(s) leave less than 1.000 uL transfection diluent. Examples: {summarize_list(low_transfection_diluent['Well'].tolist(), 12)}."
        )
    tiny_dna_warning = _tiny_dna_warning(details)
    if tiny_dna_warning:
        warnings.append(tiny_dna_warning)

    bulk = {}
    if config.get("mode") == "two_tube":
        if len(mass_groups) > 5:
            raise MixMapError(
                "L2000 supports at most 5 different total DNA masses per plate",
                details=[f"This plate has {len(mass_groups)} different positive total DNA masses."] + [
                    f"{group['Total DNA_ng']:g} ng: {summarize_list(group['Wells'])}"
                    for group in mass_groups
                ],
                fixes=["Use at most five total DNA masses per plate, or split this layout into separate plates."],
            )
        summary["Bulk transfectant mix"] = ""
        summary["Bulk mix symbol"] = ""
        well_indices = {row["Well"]: index for index, row in summary.iterrows()}
        mixes = []
        for mass_group in mass_groups:
            indices = [well_indices[well] for well in mass_group["Wells"]]
            group = summary.loc[indices]
            label = f"Bulk transfectant mix {mass_group['Number']}"
            summary.loc[indices, "Bulk transfectant mix"] = label
            summary.loc[indices, "Bulk mix symbol"] = mass_group["Symbol"]
            bulk_overage = config.get("bulk_overage_factor", 1.0)
            reagent_total = float(group["Reagent_uL"].sum()) * bulk_overage
            diluent_total = float(group["Transfection diluent_uL"].sum()) * bulk_overage
            mix = {
                "Mix": label,
                "Symbol": mass_group["Symbol"],
                "Total DNA_ng": float(mass_group["Total DNA_ng"]),
                "Wells": group["Well"].tolist(),
                "Nonempty wells": len(indices),
                "Well overage factor": config["well_overage_factor"],
                "Bulk overage factor": bulk_overage,
                "Bulk reagent_uL": reagent_total,
                "Bulk diluent_uL": diluent_total,
                "Bulk total_uL": reagent_total + diluent_total,
                "Per-well DNA mix_uL": config["final_volume_ul"] * config["well_overage_factor"] / 2.0,
                "Per-well transfection mix_uL": config["final_volume_ul"] * config["well_overage_factor"] / 2.0,
            }
            if not all(math.isfinite(value) for key, value in mix.items() if key.endswith("_uL")):
                raise MixMapError("The inputs produce non-finite bulk volumes",
                                  fixes=["Check the reagent settings for extreme values."])
            mixes.append(mix)
            if 0 < reagent_total < 1.0:
                warnings.append(
                    f"{label}: {config['reagent_label']} total is below 1.000 uL; consider increasing bulk overage or preparing a larger master mix."
                )
        bulk = {"mixes": mixes}

    mix_strings = []
    for _, row in summary.iterrows():
        if row["Total DNA_ng"] <= 0:
            mix_strings.append("")
            continue
        subset = details[details["Well"] == row["Well"]].sort_values("Slot")
        parts = [
            f"{r['Matched plasmid']}: {r['Working DNA volume_uL']:.3f}"
            for _, r in subset.iterrows()
        ]
        parts.append(f"{config['diluent_label']}: {max(row['DNA diluent_uL'], 0.0):.3f}")
        if config.get("mode") == "single_tube":
            parts.append(f"{config['reagent_label']}: {row['Reagent_uL']:.3f}")
        else:
            parts.append(f"Add {row['Transfection mix target_uL']:.3f} uL {row['Bulk transfectant mix']}")
        mix_strings.append("\n".join(parts))
    summary["Mix"] = mix_strings

    return details, summary, bulk, warnings


def safe_sheet_title(value):
    title = re.sub(r"[\[\]\:\*\?\/\\]", "_", str(value))[:31]
    return title or "Sheet"


def format_volume_ul(value):
    """Pipette increments for printable recipes, without changing calculations.

    Select precision using the original volume, including when rounding crosses
    a band boundary. Decimal half-up rounding avoids binary float tie behavior.
    Preserve the existing treatment of missing or negligible negative volumes.
    """
    value = max(Decimal(0), Decimal(str(value if value is not None else 0)))
    if value < Decimal("2.5"):
        step, places = Decimal("0.002"), 3
    elif value < 10:
        step, places = Decimal("0.01"), 2
    elif value < 20:
        step, places = Decimal("0.02"), 2
    elif value <= 200:
        step, places = Decimal("0.2"), 1
    else:
        step, places = Decimal(1), 0
    rounded = (value / step).to_integral_value(rounding=ROUND_HALF_UP) * step
    return f"{rounded:.{places}f} uL"


def _printable_map_warnings(warnings, details):
    """Round diagnostic examples from source values only in the printable map.

    Match the complete generated diagnostic (optionally reagent-prefixed), so
    numbers inside user-supplied plasmid names and other warnings stay literal.
    """
    diagnostic = _tiny_dna_warning(details)
    printable = _tiny_dna_warning(details, format_volume_ul) if diagnostic else None
    for warning in dedupe_messages(warnings or []):
        if diagnostic and (warning == diagnostic or warning.endswith(": " + diagnostic)):
            warning = warning[:-len(diagnostic)] + printable
        yield warning


def shorten_label(value, max_chars=24):
    text = clean_cell(value)
    if len(text) <= max_chars:
        return text
    return text[: max_chars - 3].rstrip() + "..."


def number_format_for_column(col_name):
    normalized = normalize_header(col_name)
    if "ul" in normalized or "volume" in normalized or "concentration" in normalized:
        return "0.000"
    if "ng" in normalized or "ratio" in normalized or "factor" in normalized:
        return "0.0"
    if "row" in normalized or "slot" in normalized:
        return "0"
    return None


def literal_cell(ws, row, column, value):
    """Keep source labels literal, including strings beginning with '='."""
    cell = ws.cell(row, column, value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def write_dataframe(ws, df, start_row=1, start_col=1, freeze=True):
    header_fill = PatternFill("solid", fgColor="D4D4D4")
    thin = Side(border_style="thin", color="D7D7D7")
    header_font = Font(name="Arial", bold=True, size=10)
    body_font = Font(name="Arial", size=10)

    for col_idx, col_name in enumerate(df.columns, start=start_col):
        cell = literal_cell(ws, start_row, col_idx, col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for row_idx, (_, row) in enumerate(df.iterrows(), start=start_row + 1):
        for col_idx, col_name in enumerate(df.columns, start=start_col):
            value = row[col_name]
            if value is not None and pd.isna(value):
                value = None
            cell = literal_cell(ws, row_idx, col_idx, value)
            cell.font = body_font
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            number_format = number_format_for_column(col_name)
            if number_format and isinstance(value, (int, float)):
                cell.number_format = number_format

    if freeze:
        ws.freeze_panes = ws.cell(start_row + 1, start_col).coordinate
    if len(df.columns):
        end_col = start_col + len(df.columns) - 1
        end_row = start_row + max(len(df), 1)
        ws.auto_filter.ref = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}"

    for col_idx, col_name in enumerate(df.columns, start=start_col):
        max_len = len(str(col_name))
        for row_idx in range(start_row + 1, start_row + len(df) + 1):
            value = ws.cell(row_idx, col_idx).value
            if value is not None:
                max_len = max(max_len, min(60, max(len(part) for part in str(value).split("\n"))))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 42)

    ws.sheet_view.showGridLines = False


def write_merged_note(ws, row_idx, last_col, text, fill, font, alignment=None):
    if last_col > 1:
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=last_col)
    cell = literal_cell(ws, row_idx, 1, text)
    cell.fill = fill
    cell.font = font
    cell.alignment = alignment or Alignment(wrap_text=True, vertical="center")
    return cell


def plate_border(row_pos, col_pos, row_count, col_count):
    thin = Side(border_style="thin", color="C8C8C8")
    medium = Side(border_style="medium", color="4F4F4F")
    row_group = 2 if row_count >= 6 else max(row_count, 1)
    col_group = 3 if col_count >= 9 else (2 if col_count >= 6 else max(col_count, 1))
    return Border(
        left=medium if col_pos == 1 or (col_pos - 1) % col_group == 0 else thin,
        right=medium if col_pos == col_count or col_pos % col_group == 0 else thin,
        top=medium if row_pos == 1 or (row_pos - 1) % row_group == 0 else thin,
        bottom=medium if row_pos == row_count or row_pos % row_group == 0 else thin,
    )


def format_well_mix_lines(row, details, config):
    if row is None or float(row["Total DNA_ng"]) <= 0:
        return []
    subset = details[details["Well"] == row["Well"]].sort_values("Slot")
    mode_label = "DNA mix" if config.get("mode") == "two_tube" else "complete mix"
    lines = [f"{row['Well']}  {mode_label}"]
    for _, plasmid_row in subset.iterrows():
        lines.append(
            f"{format_volume_ul(plasmid_row['Working DNA volume_uL'])}  "
            f"{shorten_label(plasmid_row['Matched plasmid'])}"
        )
    lines.append(f"{format_volume_ul(row['DNA diluent_uL'])}  {shorten_label(config['diluent_label'])}")
    if config.get("mode") == "single_tube":
        lines.append(f"{format_volume_ul(row['Reagent_uL'])}  {shorten_label(config['reagent_label'])}")
    else:
        symbol = row.get("Bulk mix symbol", "")
        if not symbol:
            lines.append(row["Bulk transfectant mix"])
        lines.append(f"Add {format_volume_ul(row['Transfection mix target_uL'])}")
        if symbol:
            lines.append(symbol)
    return lines


def apply_print_settings(ws, last_col, last_row):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.25
    ws.page_margins.bottom = 0.25
    ws.page_margins.header = 0.15
    ws.page_margins.footer = 0.15
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:{get_column_letter(last_col)}{last_row}"


def write_bulk_mixes_sheet(wb, bulk, config):
    """Printable recipes, using the same labels and symbols as the plate map."""
    ws = wb.create_sheet("Bulk transfectant mixes")
    title_font = Font(name="Arial", bold=True, size=15)
    body_font = Font(name="Arial", size=11)
    white = PatternFill("solid", fgColor="FFFFFF")
    recipe_fill = PatternFill("solid", fgColor="EEEEEE")
    write_merged_note(ws, 1, 4, "L2000 bulk transfectant mixes", white, title_font)
    ws.row_dimensions[1].height = 26
    write_merged_note(ws, 2, 4,
                      f"Recipes include well overage x{config['well_overage_factor']:g} and bulk overage "
                      f"x{config.get('bulk_overage_factor', 1.0):g}. Use the same factors for every mix.",
                      white, body_font)
    ws.row_dimensions[2].height = 30
    write_merged_note(ws, 3, 4,
                      f"Prepare {format_volume_ul(config['final_volume_ul'] * config['well_overage_factor'])} "
                      f"combined mix per DNA-containing well; deliver {format_volume_ul(config['final_volume_ul'])} to each of those wells. "
                      "Preparation includes extra for pipetting loss.", white, body_font)
    ws.row_dimensions[3].height = 30
    row = 5
    for mix in bulk["mixes"]:
        label = f"{mix['Mix']} {mix.get('Symbol', '')}".rstrip()
        write_merged_note(ws, row, 4,
                          f"{label} | {mix['Total DNA_ng']:.12g} ng DNA/well | {mix['Nonempty wells']} wells",
                          recipe_fill, Font(name="Arial", bold=True, size=12))
        ws.row_dimensions[row].height = 25
        recipe = (
            f"Prepare: {format_volume_ul(mix['Bulk reagent_uL'])} {config['reagent_label']} + "
            f"{format_volume_ul(mix['Bulk diluent_uL'])} {config['diluent_label']} = "
            f"{format_volume_ul(mix['Bulk total_uL'])} total.\n"
            f"For each assigned well: combine {format_volume_ul(mix['Per-well DNA mix_uL'])} DNA mix "
            f"with {format_volume_ul(mix['Per-well transfection mix_uL'])} of this bulk transfectant mix."
        )
        write_merged_note(ws, row + 1, 4, recipe, white, body_font)
        ws.row_dimensions[row + 1].height = 38
        wells = "Assigned wells: " + ", ".join(mix["Wells"])
        write_merged_note(ws, row + 2, 4, wells, white, body_font)
        ws.row_dimensions[row + 2].height = max(20, 15 * math.ceil(len(wells) / 110))
        row += 4
    for column in "ABCD":
        ws.column_dimensions[column].width = 30
    apply_print_settings(ws, 4, row - 1)
    ws.freeze_panes = "A5"

    # Keep machine-readable recipe amounts on the same sheet, outside the
    # printable recipe cards. Numeric cells retain full calculation precision.
    records = [{**mix, "Wells": ", ".join(mix["Wells"])} for mix in bulk["mixes"]]
    write_dataframe(ws, pd.DataFrame(records), start_row=row + 1, freeze=False)
    # write_dataframe sizes columns for the numeric table; restore print widths.
    for column in "ABCD":
        ws.column_dimensions[column].width = 30


def write_mix_map_sheet(
    ws,
    csv_name,
    spreadsheet_id,
    config,
    details,
    summary,
    bulk,
    case_warnings,
    plate_format,
    rows,
):
    """Render a complete plate or one printable half without recalculating recipes."""
    cols = plate_format.columns
    start_col = 2
    # Excel may clip a single column wider than the printable page even with
    # fit-to-page enabled. Keep the dish as one physical well by merging eight
    # ordinary-width worksheet columns for its header and recipe cell.
    column_span = 8 if plate_format.well_count == 1 else 1
    last_col = start_col + len(cols) * column_span - 1
    # Preserve the 48-well map's width, spreading it over fewer actual wells
    # for small plates so the title and preparation notes remain printable.
    well_col_width = (190 if len(cols) == 12 else 176) / len(cols)

    title_fill = PatternFill("solid", fgColor="FFFFFF")
    header_fill = PatternFill("solid", fgColor="D4D4D4")
    note_fill = PatternFill("solid", fgColor="EEEEEE")
    active_fill = PatternFill("solid", fgColor="F7F7F7")
    empty_fill = PatternFill("solid", fgColor="E6E6E6")
    title_font = Font(name="Arial", bold=True, size=17)
    subtitle_font = Font(name="Arial", size=8, color="666666")
    note_font = Font(name="Arial", bold=True, size=9)
    header_font = Font(name="Arial", bold=True, size=11)
    well_font_size = 8 if len(cols) >= 10 else 8.5
    well_font = Font(name="Arial", size=well_font_size)
    well_label_font = InlineFont(rFont="Arial", sz=well_font_size, b=True)
    symbol_font = InlineFont(rFont="Arial", sz=16, b=True)
    empty_font = Font(name="Arial", size=8, color="666666")

    title = f"{Path(csv_name).stem} - {config['display_name']} print mix map"
    if plate_format.well_count == 96:
        title += f" - Rows {rows[0]}-{rows[-1]}"
    write_merged_note(
        ws,
        1,
        last_col,
        title,
        title_fill,
        title_font,
        Alignment(horizontal="left", vertical="center"),
    )
    ws.row_dimensions[1].height = 24
    write_merged_note(
        ws,
        2,
        last_col,
        f"{plate_format.label} ({plate_format.well_range}) | "
        f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        title_fill,
        subtitle_font,
        Alignment(horizontal="left", vertical="center"),
    )
    ws.row_dimensions[2].height = 15

    notes = []
    if config.get("mode") == "two_tube":
        assignment = "symbol" if len(bulk.get("mixes", [])) > 1 else "mix name"
        notes.append(
            "Prepare each bulk mix in a separate tube using the instructions below. "
            f"Each well shows its assigned {assignment} and the aliquot to add to its DNA mix."
        )
        notes.append(
            f"Prepare {format_volume_ul(config['final_volume_ul'] * config['well_overage_factor'] / 2)} "
            f"DNA mix + the designated bulk aliquot = "
            f"{format_volume_ul(config['final_volume_ul'] * config['well_overage_factor'])} combined mix per well."
        )
    else:
        notes.append(
            f"Prepare {format_volume_ul(config['final_volume_ul'] * config['well_overage_factor'])} "
            "complete DNA/diluent/reagent mix for each populated well using the listed components."
        )
    notes.append(f"Deliver {format_volume_ul(config['final_volume_ul'])} to each DNA-containing well. "
                 f"Preparation volumes include well overage x{config['well_overage_factor']:g} for pipetting loss.")
    if plate_format.well_count == 96 and bulk.get("mixes"):
        notes.append(
            "Whole-plate totals: prepare each bulk recipe once for the entire plate. "
            "The same recipes are repeated on both printable halves."
        )
    for mix in bulk.get("mixes", []):
        label = f"{mix['Mix']} {mix.get('Symbol', '')}".rstrip()
        notes.append(f"{label} | {mix['Total DNA_ng']:.12g} ng DNA/well | "
                     f"{mix['Nonempty wells']} wells | "
                     f"Add {format_volume_ul(mix['Bulk reagent_uL'])} {config['reagent_label']} to "
                     f"{format_volume_ul(mix['Bulk diluent_uL'])} {config['diluent_label']}. | "
                     f"Add {format_volume_ul(mix['Per-well transfection mix_uL'])} "
                     "to each assigned DNA mix.")
    for warning in limited_examples(_printable_map_warnings(case_warnings, details), 3):
        notes.append(f"WARNING: {warning}")

    for offset, note in enumerate(notes):
        row_idx = 3 + offset
        write_merged_note(ws, row_idx, last_col, note, note_fill, note_font)
        if config.get("mode") == "two_tube":
            wrapped_lines = sum(max(1, math.ceil(len(line) / 120)) for line in note.splitlines())
            ws.row_dimensions[row_idx].height = max(30 if note.startswith("WARNING:") else 18, 15 * wrapped_lines)
        else:
            ws.row_dimensions[row_idx].height = 30 if note.startswith("WARNING:") or len(note) > 120 else 18

    start_row = 3 + len(notes) + 1
    corner = ws.cell(start_row, 1)
    corner.font = header_font
    corner.fill = header_fill
    corner.alignment = Alignment(horizontal="center", vertical="center")
    corner.border = plate_border(1, 1, max(len(rows), 1), max(len(cols), 1))

    for col_pos, col_num in enumerate(cols, start=1):
        sheet_col = start_col + (col_pos - 1) * column_span
        cell = ws.cell(start_row, sheet_col, col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = plate_border(1, col_pos, max(len(rows), 1), len(cols))
        if column_span > 1:
            ws.merge_cells(start_row=start_row, start_column=sheet_col,
                           end_row=start_row, end_column=sheet_col + column_span - 1)
    ws.row_dimensions[start_row].height = 21

    summary_lookup = {(row["Row"], int(row["Col"])): row for _, row in summary.iterrows()}
    max_lines_by_row = {}
    symbol_rows = set()
    for row_pos, row_label in enumerate(rows, start=1):
        sheet_row = start_row + row_pos
        label = ws.cell(sheet_row, 1, row_label)
        label.font = header_font
        label.fill = header_fill
        label.alignment = Alignment(horizontal="center", vertical="center")
        label.border = plate_border(row_pos, 1, len(rows), max(len(cols), 1))
        max_lines = 1
        for col_pos, col_num in enumerate(cols, start=1):
            summary_row = summary_lookup.get((row_label, col_num))
            lines = format_well_mix_lines(summary_row, details, config) if summary_row is not None else []
            # Account for automatic wrapping, especially the narrower 96-well
            # columns. The estimate includes space for the cell's padding.
            chars_per_line = max(1, int((well_col_width * 7 - 8) / (well_font_size * 0.7)))
            wrapped_lines = sum(max(1, math.ceil(len(line) / chars_per_line)) for line in lines)
            max_lines = max(max_lines, wrapped_lines or 1)
            value = "\n".join(lines) if lines else ("A1" if plate_format.well_count == 1 else "")
            if lines:
                well_label = str(summary_row["Well"])
                value = CellRichText(TextBlock(well_label_font, well_label))
                has_symbol = config.get("mode") == "two_tube" and summary_row.get("Bulk mix symbol")
                body_lines = lines[:-1] if has_symbol else lines
                value.append("\n".join(body_lines)[len(well_label):])
                if has_symbol:
                    value.append("\n")
                    value.append(TextBlock(symbol_font, lines[-1]))
                    symbol_rows.add(row_label)
            sheet_col = start_col + (col_pos - 1) * column_span
            cell = ws.cell(sheet_row, sheet_col, value)
            cell.alignment = Alignment(wrap_text=True, horizontal="left", vertical="top")
            cell.border = plate_border(row_pos, col_pos, len(rows), len(cols))
            if lines:
                cell.fill = active_fill
                cell.font = well_font
            else:
                cell.fill = empty_fill
                cell.font = empty_font
            if column_span > 1:
                ws.merge_cells(start_row=sheet_row, start_column=sheet_col,
                               end_row=sheet_row, end_column=sheet_col + column_span - 1)
        max_lines_by_row[row_label] = max_lines

    ws.column_dimensions["A"].width = 4.5
    for j in range(start_col, last_col + 1):
        ws.column_dimensions[get_column_letter(j)].width = well_col_width / column_span
    for row_pos, row_label in enumerate(rows, start=1):
        sheet_row = start_row + row_pos
        line_count = max_lines_by_row.get(row_label, 1)
        # Reserve a taller final line for the print-friendly mix symbol.
        symbol_extra_height = 10 if row_label in symbol_rows else 0
        ws.row_dimensions[sheet_row].height = max(48, 14 + 12 * line_count + symbol_extra_height)

    last_grid_row = start_row + len(rows)
    ws.freeze_panes = ws.cell(start_row + 1, start_col).coordinate
    apply_print_settings(ws, last_col, last_grid_row)


def write_mix_map_workbook(
    output_path,
    csv_name,
    spreadsheet_id,
    config_name,
    config,
    details,
    summary,
    bulk,
    concentrations_df,
    case_warnings=None,
    *,
    plate_type=48,
):
    plate_format = get_plate_format(plate_type)
    wb = Workbook()
    for index, (sheet_name, rows) in enumerate(plate_format.map_sections):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = sheet_name
        write_mix_map_sheet(
            ws, csv_name, spreadsheet_id, config, details, summary, bulk,
            case_warnings, plate_format, rows,
        )

    if bulk:
        write_bulk_mixes_sheet(wb, bulk, config)

    details_ws = wb.create_sheet("Per-well details")
    detail_cols = [
        "Well",
        "Slot",
        "Matched plasmid",
        "Mass_ng",
        "Concentration_ng_per_uL",
        "Base DNA volume_uL",
        "Working DNA volume_uL",
        "Concentration source worksheet",
        "Concentration source row",
        "CSV row",
    ]
    detail_out = details[detail_cols].copy() if not details.empty else pd.DataFrame(columns=detail_cols)
    write_dataframe(details_ws, detail_out)

    used_names = set(details["Matched plasmid"]) if not details.empty else set()
    conc_ws = wb.create_sheet("Concentrations used")
    used_conc = concentrations_df[concentrations_df["Plasmid"].isin(used_names)].copy()
    write_dataframe(conc_ws, used_conc)

    summary_ws = wb.create_sheet("Well summary")
    summary_cols = [
        "Well",
        "Total DNA_ng",
        "Total working DNA_uL",
        "DNA mix target_uL",
        "DNA diluent_uL",
        "Reagent_uL",
        "Transfection mix target_uL",
        "Transfection diluent_uL",
    ]
    if bulk:
        summary_cols.extend(["Bulk transfectant mix", "Bulk mix symbol"])
    write_dataframe(summary_ws, summary[summary_cols].copy())

    if case_warnings:
        warnings_ws = wb.create_sheet("Warnings")
        warnings_out = pd.DataFrame(
            [
                {
                    "Severity": "Warning",
                    "Message": warning,
                    "Suggested action": "Review the source CSV, plasmid concentration table, and run settings before pipetting.",
                }
                for warning in dedupe_messages(case_warnings)
            ]
        )
        write_dataframe(warnings_ws, warnings_out)

    config_ws = wb.create_sheet("Run config")
    config_rows = [
        {"Setting": "CSV file", "Value": csv_name},
        {"Setting": "Concentration source", "Value": spreadsheet_id},
        {"Setting": "Reagent config", "Value": config_name},
        {"Setting": "Display name", "Value": config["display_name"]},
        {"Setting": "Mode", "Value": config.get("mode")},
        {"Setting": "Final volume_uL", "Value": config["final_volume_ul"]},
        {"Setting": "Ratio_uL_per_ug", "Value": config["dna_to_reagent_ratio_ul_per_ug"]},
        {"Setting": "Well overage factor", "Value": config["well_overage_factor"]},
        {"Setting": "Bulk overage factor", "Value": config.get("bulk_overage_factor", 1.0)},
        {"Setting": "Plate type", "Value": plate_format.label},
        {"Setting": "Plate wells", "Value": plate_format.well_count},
        {"Setting": "Plate rows", "Value": ", ".join(plate_format.rows)},
        {"Setting": "Plate columns", "Value": len(plate_format.columns)},
        {"Setting": "Plate well range", "Value": plate_format.well_range},
        {"Setting": "Print layout", "Value": (
            "Letter landscape, one page per half (A-D and E-H)"
            if plate_format.well_count == 96 else "Letter landscape, fit to one page"
        )},
        {"Setting": "Map cell format", "Value": "Well label first, then volume-first pipetting lines"},
    ]
    for warning in case_warnings or []:
        config_rows.append({"Setting": "Warning", "Value": warning})
    write_dataframe(config_ws, pd.DataFrame(config_rows))

    wb.active = 0
    if isinstance(output_path, (str, bytes, os.PathLike)):
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    wb.save(output_path)
    return output_path
