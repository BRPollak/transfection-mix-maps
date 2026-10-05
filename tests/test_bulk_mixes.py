"""Regressions for separate L2000 recipes and unambiguous bench assignments."""

import csv
from decimal import Decimal
from io import BytesIO, StringIO
import re

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
import pandas as pd
import pytest

import core
from sources import read_plate
from workflow import default_configs, generate, generate_batch


STOCKS = {"Stocks": [["Plasmid", "Concentration"], ["A", "100"], ["B", "100"]]}


def plate_bytes(rows, wide=False):
    stream = StringIO()
    writer = csv.writer(stream)
    writer.writerow(["Well", "Plasmid1" if wide else "Plasmid", "Mass1 (ng)" if wide else "Mass (ng)"])
    writer.writerows(rows)
    return stream.getvalue().encode()


def calculate(rows, config=None, wide=False):
    wells, long, _ = core.standardize_plate_csv(read_plate(plate_bytes(rows, wide=wide)))
    _, lookup, _ = core.load_plasmid_concentrations(STOCKS, used_plasmids=long.Plasmid.tolist())
    matched, _ = core.match_concentrations(long, lookup)
    return core.calculate_mix(wells, matched, config or default_configs()["L2000"])


def sheet_records(sheet):
    rows = sheet.iter_rows(values_only=True)
    if sheet.title == "Bulk transfectant mixes":
        # Printable recipe cards precede the full-precision numeric table.
        headers = next(row for row in rows if "Mix" in row and "Total DNA_ng" in row)
    else:
        headers = next(rows)
    return [dict(zip(headers, row)) for row in rows]


def well_cell(sheet, well):
    matches = [cell for row in sheet for cell in row
               if isinstance(cell.value, str) and cell.value.split()[:1] == [well]]
    assert len(matches) == 1, f"Expected exactly one plate cell for {well}"
    return matches[0]


def mix_legend_cell(sheet, label, symbol=""):
    heading = f"{label} {symbol}".rstrip()
    matches = [cell for row in sheet for cell in row
               if isinstance(cell.value, str) and cell.value.startswith(heading + " |")]
    assert len(matches) == 1, f"Expected exactly one legend or recipe card for {label}"
    return matches[0]


@pytest.mark.parametrize("entries", [None, pd.DataFrame()])
def test_mass_grouping_without_entries_is_empty(entries):
    assert core.group_dna_masses(entries) == []


def test_shared_mass_groups_normalize_wells_and_preserve_exact_totals():
    entries = pd.DataFrame([
        {"Well": "a01", "Mass_ng": "0.1"}, {"Well": "A1", "Mass_ng": "0.2"},
        {"Well": "A2", "Mass_ng": "0.3"}, {"Well": "A3", "Mass_ng": "0.3000001"},
        {"Well": "G1", "Mass_ng": "0.3000001"}, {"Well": "G2", "Mass_ng": "0.3000001"},
    ])
    groups = core.group_dna_masses(entries)
    assert groups == [
        {"Number": 1, "Total DNA_ng": Decimal("0.3000001"), "Wells": ["A3", "G1", "G2"], "Symbol": "▲"},
        {"Number": 2, "Total DNA_ng": Decimal("0.3"), "Wells": ["A1", "A2"], "Symbol": "●"},
    ]
    assert core.group_dna_masses(entries.iloc[::-1]) == groups


def test_mass_grouping_ignores_nonpositive_and_nonfinite_masses():
    entries = pd.DataFrame([{"Well": f"A{number}", "Mass_ng": mass}
                            for number, mass in enumerate([100, 0, -100, float("nan"), float("inf"), None], 1)])
    assert core.group_dna_masses(entries) == [
        {"Number": 1, "Total DNA_ng": Decimal(100), "Wells": ["A1"], "Symbol": ""},
    ]


def test_more_than_five_mass_groups_get_distinct_number_markers():
    entries = pd.DataFrame([{"Well": f"A{number}", "Mass_ng": number * 100} for number in range(1, 7)])
    assert [group["Symbol"] for group in core.group_dna_masses(entries)] == [str(number) for number in range(1, 7)]


def test_different_dna_totals_receive_the_correct_l2000_concentration():
    _, summary, bulk, _ = calculate([("A1", "A", 100), ("A2", "A", 300)])
    mixes = bulk["mixes"]
    assert len(mixes) == 2
    by_well = summary.set_index("Well")
    # The two 15-uL aliquots must contain 0.24 and 0.72 uL reagent,
    # respectively. A shared tube would incorrectly give both 0.48 uL.
    for mix, well, mass, reagent, diluent in zip(
            mixes, ["A1", "A2"], [100, 300], [0.24, 0.72], [14.76, 14.28]):
        assert mix["Total DNA_ng"] == mass
        assert mix["Wells"] == [well]
        assert mix["Nonempty wells"] == 1
        assert mix["Bulk reagent_uL"] == pytest.approx(reagent * 1.2)
        assert mix["Bulk diluent_uL"] == pytest.approx(diluent * 1.2)
        assert mix["Bulk total_uL"] == pytest.approx(18)
        assert mix["Per-well DNA mix_uL"] == pytest.approx(15)
        assert mix["Per-well transfection mix_uL"] == pytest.approx(15)
        aliquot_reagent = mix["Bulk reagent_uL"] / mix["Bulk total_uL"] * mix["Per-well transfection mix_uL"]
        assert aliquot_reagent == pytest.approx(reagent)
        assert aliquot_reagent == pytest.approx(by_well.loc[well, "Reagent_uL"])
        assert by_well.loc[well, "Bulk transfectant mix"] == mix["Mix"]
        assert by_well.loc[well, "Bulk mix symbol"] == mix["Symbol"]


def test_groups_rank_by_well_count_then_physical_position_independent_of_input_order():
    rows = [("AA1", "A", 400), ("Z1", "A", 300),
            ("A10", "A", 100), ("B1", "A", 100),
            ("a02", "A", 200), ("B2", "A", 200),
            ("C3", "A", 500), ("C1", "A", 500), ("C2", "A", 500)]
    expected_wells = [["C1", "C2", "C3"], ["A2", "B2"], ["A10", "B1"], ["Z1"], ["AA1"]]
    _, summary, bulk, _ = calculate(rows)
    assert [mix["Wells"] for mix in bulk["mixes"]] == expected_wells
    assert [mix["Total DNA_ng"] for mix in bulk["mixes"]] == [500, 200, 100, 300, 400]
    assert [mix["Mix"] for mix in bulk["mixes"]] == [f"Bulk transfectant mix {n}" for n in range(1, 6)]
    assert [mix["Symbol"] for mix in bulk["mixes"]] == list(core.BULK_MIX_SYMBOLS)
    _, reversed_summary, reversed_bulk, _ = calculate(list(reversed(rows)))
    assert reversed_bulk == bulk
    assert reversed_summary.set_index("Well")["Bulk transfectant mix"].to_dict() == summary.set_index("Well")["Bulk transfectant mix"].to_dict()


@pytest.mark.parametrize("count", [1, 5])
def test_accepts_one_through_five_positive_dna_groups(count):
    _, summary, bulk, _ = calculate([(f"A{n}", "A", n * 100) for n in range(1, count + 1)])
    assert len(bulk["mixes"]) == count
    assert summary["Bulk transfectant mix"].tolist() == [f"Bulk transfectant mix {n}" for n in range(1, count + 1)]


def test_six_positive_dna_groups_are_rejected():
    with pytest.raises(core.MixMapError, match=r"(?i)(5|five)"):
        calculate([(f"A{n}", "A", n * 100) for n in range(1, 7)])


def test_empty_and_zero_mass_wells_do_not_consume_a_group_or_reagent():
    rows = [(f"A{n}", "A", n * 100) for n in range(1, 6)] + [("A6", "A", 0), ("A7", "", "")]
    _, summary, bulk, _ = calculate(rows, wide=True)
    assert len(bulk["mixes"]) == 5
    assert sum(mix["Nonempty wells"] for mix in bulk["mixes"]) == 5
    assert sum(mix["Bulk total_uL"] for mix in bulk["mixes"]) == pytest.approx(90)
    for well in ["A6", "A7"]:
        assert summary.set_index("Well").loc[well, "Bulk transfectant mix"] == ""
        assert summary.set_index("Well").loc[well, "Bulk mix symbol"] == ""
        assert all(well not in mix["Wells"] for mix in bulk["mixes"])


def test_identical_totals_share_a_group_despite_different_plasmid_compositions():
    _, _, bulk, _ = calculate([("A1", "A", 75), ("A1", "B", 25), ("B1", "A", 100)])
    assert len(bulk["mixes"]) == 1
    assert bulk["mixes"][0]["Wells"] == ["A1", "B1"]
    assert bulk["mixes"][0]["Total DNA_ng"] == 100


def test_equal_decimal_totals_share_a_group_without_merging_distinct_close_totals():
    rows = [("A1", "A", "0.1"), ("A1", "B", "0.2"), ("A2", "A", "0.3"),
            ("A3", "A", "0.3000001"), ("A4", "A", "0.3000002")]
    _, _, bulk, _ = calculate(rows)
    assert [mix["Wells"] for mix in bulk["mixes"]] == [["A1", "A2"], ["A3"], ["A4"]]
    assert [mix["Total DNA_ng"] for mix in bulk["mixes"]] == pytest.approx([0.3, 0.3000001, 0.3000002])


def test_custom_overages_are_applied_identically_to_every_group():
    config = default_configs()["L2000"]
    config.update(final_volume_ul=40, dna_to_reagent_ratio_ul_per_ug=3,
                  well_overage_factor=1.5, bulk_overage_factor=1.7)
    _, summary, bulk, _ = calculate([("A1", "A", 100), ("A2", "A", 100), ("B1", "A", 300)], config)
    for mix, count, per_well_reagent in zip(bulk["mixes"], [2, 1], [0.45, 1.35]):
        assert mix["Well overage factor"] == 1.5
        assert mix["Bulk overage factor"] == 1.7
        assert mix["Nonempty wells"] == count
        assert mix["Per-well DNA mix_uL"] == pytest.approx(30)
        assert mix["Per-well transfection mix_uL"] == pytest.approx(30)
        assert mix["Bulk reagent_uL"] == pytest.approx(per_well_reagent * count * 1.7)
        assert mix["Bulk diluent_uL"] == pytest.approx((30 - per_well_reagent) * count * 1.7)
        assert mix["Bulk total_uL"] == pytest.approx(30 * count * 1.7)
    assert summary["DNA mix target_uL"].tolist() == [30, 30, 30]


def test_six_groups_in_one_plate_prevent_every_workbook_in_the_batch(tmp_path, monkeypatch):
    plates = [{"name": "valid.csv", "bytes": plate_bytes([("A1", "A", 100)])},
              {"name": "six-groups.csv", "bytes": plate_bytes([(f"A{n}", "A", n * 100) for n in range(1, 7)])}]

    def unexpected_write(*args, **kwargs):
        pytest.fail("No workbook should be staged before every plate passes validation")

    monkeypatch.setattr(core, "write_mix_map_workbook", unexpected_write)
    with pytest.raises(core.MixMapError, match="no output was made") as caught:
        generate_batch(plates, STOCKS, "test", "test-time", ["L2000"], default_configs())
    assert any("six-groups.csv" in detail for detail in caught.value.details)
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("group_count", [2, 5])
def test_mix_map_and_recipe_sheet_identify_every_wells_recipe_and_aliquot(tmp_path, group_count):
    rows = [(f"A{n}", "A", n * 100) for n in range(1, group_count + 1)] + [("B1", "B", 100)]
    config = default_configs()["L2000"]
    result = generate(plate_bytes(rows), "mixed.csv", STOCKS, "test", "test-time",
                      ["L2000"], default_configs())
    artifact = result["artifacts"][0]
    with BytesIO(artifact["bytes"]) as stream:
        workbook = load_workbook(stream)
        try:
            recipe_rows = sheet_records(workbook["Bulk transfectant mixes"])
            assert len(recipe_rows) == group_count
            summary = {row["Well"]: row for row in sheet_records(workbook["Well summary"])}
            group_symbols, legend_fills, card_fills, well_fills = set(), set(), set(), set()
            for recipe, mix in zip(recipe_rows, artifact["bulk"]["mixes"]):
                assert recipe["Mix"] == mix["Mix"]
                assert recipe["Symbol"] == mix["Symbol"]
                assert recipe["Total DNA_ng"] == mix["Total DNA_ng"]
                legend = mix_legend_cell(workbook["Mix Map"], mix["Mix"], mix["Symbol"])
                card = mix_legend_cell(workbook["Bulk transfectant mixes"], mix["Mix"], mix["Symbol"])
                group_symbols.add(mix["Symbol"])
                legend_fills.add(legend.fill.fgColor.rgb)
                card_fills.add(card.fill.fgColor.rgb)
                assert (f"Add {core.format_volume_ul(mix['Bulk reagent_uL'])} {config['reagent_label']} to "
                        f"{core.format_volume_ul(mix['Bulk diluent_uL'])} {config['diluent_label']}.") in legend.value
                assert "Add 15.00 uL to each assigned DNA mix." in legend.value
                assert legend.alignment.wrap_text
                assert workbook["Mix Map"].row_dimensions[legend.row].height >= 30
                for field in ["Bulk reagent_uL", "Bulk diluent_uL", "Bulk total_uL", "Per-well transfection mix_uL"]:
                    assert recipe[field] == pytest.approx(mix[field])
                for well in mix["Wells"]:
                    assert well in str(recipe["Wells"])
                    cell = well_cell(workbook["Mix Map"], well)
                    assert mix["Mix"] not in cell.value
                    assert cell.value.splitlines()[-1] == mix["Symbol"]
                    assert re.search(r"Add\s+15(?:\.0+)?\s*(?:µL|uL)", cell.value)
                    assert cell.alignment.wrap_text
                    well_fills.add(cell.fill.fgColor.rgb)
                    assert summary[well]["Bulk transfectant mix"] == mix["Mix"]
                    assert summary[well]["Bulk mix symbol"] == mix["Symbol"]
            assert len(group_symbols) == group_count
            assert legend_fills == card_fills == {"00EEEEEE"}
            assert well_fills == {"00F7F7F7"}
        finally:
            workbook.close()


def test_each_plate_restarts_mix_numbering_and_has_its_own_recipes(tmp_path):
    plates = [{"name": "one.csv", "bytes": plate_bytes([("A1", "A", 100), ("A2", "A", 300)])},
              {"name": "two.csv", "bytes": plate_bytes([("B1", "A", 500)])}]
    result = generate_batch(plates, STOCKS, "test", "test-time", ["L2000"], default_configs())
    assert len(result["artifacts"]) == 2
    for artifact, expected_masses in zip(result["artifacts"], [[100, 300], [500]]):
        mixes = artifact["bulk"]["mixes"]
        assert [mix["Mix"] for mix in mixes] == [f"Bulk transfectant mix {n}" for n in range(1, len(mixes) + 1)]
        assert [mix["Total DNA_ng"] for mix in mixes] == expected_masses
        assert [mix["Symbol"] for mix in mixes] == (["▲", "●"] if len(mixes) > 1 else [""])
        workbook = load_workbook(BytesIO(artifact["bytes"]))
        try:
            assert len(sheet_records(workbook["Bulk transfectant mixes"])) == len(mixes)
            if len(mixes) == 1:
                assert mixes[0]["Mix"] in well_cell(workbook["Mix Map"], "B1").value
                assert not any(symbol in cell.value for sheet in workbook for row in sheet for cell in row
                               if isinstance(cell.value, str) for symbol in core.BULK_MIX_SYMBOLS)
        finally:
            workbook.close()


@pytest.mark.parametrize("group_count", [1, 2])
@pytest.mark.parametrize("plate_type", [96, 48, 24, 12, 6])
def test_bold_well_labels_preserve_large_multi_mix_markers(tmp_path, group_count, plate_type):
    rows = [(f"A{number}", "A", number * 100) for number in range(1, group_count + 1)]
    result = generate(plate_bytes(rows), "markers.csv", STOCKS, "test", "test-time",
                      ["L2000"], default_configs(), plate_type=plate_type)
    data = result["artifacts"][0]["bytes"]
    plain = load_workbook(BytesIO(data))
    rich = load_workbook(BytesIO(data), rich_text=True)
    map_name = "Mix Map (A-D)" if plate_type == 96 else "Mix Map"
    try:
        for mix in result["artifacts"][0]["bulk"]["mixes"]:
            for well in mix["Wells"]:
                cell = well_cell(plain[map_name], well)
                styled = rich[map_name][cell.coordinate]
                assert str(styled.value) == cell.value
                assert styled.font.name == "Arial"
                assert styled.font.sz == (8 if plate_type == 96 else 8.5)
                assert isinstance(styled.value, CellRichText)
                label = styled.value[0]
                assert isinstance(label, TextBlock)
                assert label.text == well
                assert label.font.b is True
                assert not styled.font.b
                bold_texts = [part.text for part in styled.value if isinstance(part, TextBlock) and part.font.b]
                assert bold_texts == ([well] if group_count == 1 else [well, mix["Symbol"]])
                if group_count == 1:
                    continue
                marker = styled.value[-1]
                assert isinstance(marker, TextBlock)
                assert marker.text == mix["Symbol"]
                assert marker.font.rFont == "Arial"
                assert marker.font.sz == 16
                assert marker.font.b is True
                body_lines = len(cell.value.splitlines()) - 1
                assert rich[map_name].row_dimensions[cell.row].height >= 14 + 12 * body_lines + 20
    finally:
        plain.close()
        rich.close()


def test_lt1_has_no_group_limit_or_bulk_recipes(tmp_path):
    rows = [(f"A{n}", "A", n * 100) for n in range(1, 7)]
    result = generate(plate_bytes(rows), "lt1.csv", STOCKS, "test", "test-time",
                      ["LT1"], default_configs())
    artifact = result["artifacts"][0]
    assert artifact["bulk"] == {}
    workbook = load_workbook(BytesIO(artifact["bytes"]))
    try:
        assert "Bulk transfectant mixes" not in workbook.sheetnames
        assert "Bulk mix symbol" not in [cell.value for cell in workbook["Well summary"][1]]
        assert not any(symbol in cell.value for row in workbook["Mix Map"] for cell in row
                       if isinstance(cell.value, str) for symbol in core.BULK_MIX_SYMBOLS)
    finally:
        workbook.close()
