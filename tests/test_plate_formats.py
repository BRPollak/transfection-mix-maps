"""Selected plate geometry must constrain DNA, while tolerating unused templates."""
from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

import core
from plate_formats import get_plate_format
from workflow import default_configs, generate, generate_batch


FORMATS = [(96, "H", 12), (48, "F", 8), (24, "D", 6), (12, "C", 4), (6, "B", 3), (1, "A", 1)]
STOCKS = {"Stocks": [["Plasmid", "Concentration"], ["Stock", "100"]]}


def frame(rows, wide=False):
    headers = ["Well", "Plasmid1", "Mass1 (ng)"] if wide else ["Well", "Plasmid", "Mass (ng)"]
    return pd.DataFrame(rows, columns=headers)


@pytest.mark.parametrize("plate_type,last_row,last_column", FORMATS)
@pytest.mark.parametrize("wide", [False, True])
def test_each_format_accepts_first_and_last_wells(plate_type, last_row, last_column, wide):
    last = f"{last_row}{last_column}"
    rows = [["A1", "Stock", 100]]
    if last != "A1":
        rows.append([f" {last_row.lower()} 0{last_column} ", "Stock", 200])
    wells, long, warnings = core.standardize_plate_csv(frame(rows, wide), plate_type=plate_type)
    assert wells.Well.tolist() == list(dict.fromkeys(["A1", last]))
    assert long.Well.tolist() == wells.Well.tolist()
    assert not warnings


@pytest.mark.parametrize("plate_type,last_row,last_column", FORMATS)
@pytest.mark.parametrize("outside_axis", ["row", "column"])
@pytest.mark.parametrize("wide", [False, True])
def test_positive_mass_outside_each_boundary_reports_location(plate_type, last_row, last_column,
                                                             outside_axis, wide):
    outside = f"{chr(ord(last_row) + 1)}1" if outside_axis == "row" else f"A{last_column + 1}"
    with pytest.raises(core.MixMapError, match="outside the selected plate type") as caught:
        core.standardize_plate_csv(frame([["A1", "Stock", 100], [outside, "Stock", 1]], wide),
                                   plate_type=plate_type)
    detail = caught.value.details[0]
    assert f"CSV row 3, well {outside}" in detail
    assert get_plate_format(plate_type).well_range in detail


@pytest.mark.parametrize("plate_type,last_row,last_column", FORMATS)
@pytest.mark.parametrize("wide", [False, True])
def test_unused_template_wells_and_their_duplicates_are_ignored(plate_type, last_row, last_column, wide):
    outside_row = f"{chr(ord(last_row) + 1)}1"
    outside_column = f"A{last_column + 1}"
    rows = [["A1", "Stock", 100], [outside_row, "", ""], [outside_row, "", ""],
            [outside_column, "Stock", "0 ng"], [outside_column, "Stock", "+0e4"],
            ["ZZ999", "", ""], ["ZZ999", "Stock", "-0"]]
    wells, long, warnings = core.standardize_plate_csv(frame(rows, wide), plate_type=plate_type)
    assert wells.to_dict("records") == [{"Well": "A1", "Row": "A", "Col": 1}]
    assert long.Well.tolist() == ["A1"]
    assert not warnings


@pytest.mark.parametrize("wide", [False, True])
def test_empty_supported_wells_remain_part_of_layout(wide):
    wells, long, warnings = core.standardize_plate_csv(
        frame([["A1", "Stock", 100], ["B3", "", ""], ["A3", "Stock", 0], ["C8", "", ""]], wide),
        plate_type=6)
    assert wells.Well.tolist() == ["A1", "B3", "A3"]
    assert long.Well.tolist() == ["A1"]
    assert len(warnings) == 1 and "well A3" in warnings[0]


@pytest.mark.parametrize("second_entry", [["B3", "", ""], ["B3", "Stock", 0], ["B3", "Stock", 5]])
def test_supported_duplicate_wide_wells_still_fail(second_entry):
    with pytest.raises(core.MixMapError, match="Duplicate well"):
        core.standardize_plate_csv(
            frame([["A1", "Stock", 100], ["b03", "", ""], second_entry], wide=True), plate_type=6)


@pytest.mark.parametrize("wide", [False, True])
def test_positive_unsupported_duplicates_report_dna_bounds_instead_of_duplicate_wells(wide):
    with pytest.raises(core.MixMapError, match="outside the selected plate type") as caught:
        core.standardize_plate_csv(
            frame([["C8", "", ""], ["c08", "Stock", 100], ["C8", "Stock", 50]], wide), plate_type=6)
    assert len(caught.value.details) == 2
    assert all("well C8" in detail for detail in caught.value.details)


@pytest.mark.parametrize("wide", [False, True])
@pytest.mark.parametrize("plasmid,mass,reason", [
    ("Stock", "", "unreadable DNA mass"), ("Stock", "typo", "unreadable DNA mass"),
    ("Stock", "1e999", "unreadable DNA mass"), ("Stock", "-1", "Negative DNA mass"),
    ("", "100", "plasmid name"), ("", "0", "plasmid name"),
])
def test_invalid_unsupported_entries_still_fail(wide, plasmid, mass, reason):
    with pytest.raises(core.MixMapError, match=reason) as caught:
        core.standardize_plate_csv(frame([["A1", "Stock", 100], ["C8", plasmid, mass]], wide), plate_type=6)
    assert "CSV row 3" in caught.value.details[0]


@pytest.mark.parametrize("wide", [False, True])
def test_malformed_coordinates_still_fail_for_blank_template_rows(wide):
    with pytest.raises(core.MixMapError, match="Malformed well"):
        core.standardize_plate_csv(frame([["A1", "Stock", 100], ["C0", "", ""]], wide), plate_type=6)


@pytest.mark.parametrize("wide", [False, True])
@pytest.mark.parametrize("well", ["A1", "C8"])
def test_entirely_blank_or_zero_layouts_cannot_generate(wide, well):
    with pytest.raises(core.MixMapError, match="No usable plasmid/mass"):
        core.standardize_plate_csv(frame([[well, "", ""]], wide), plate_type=6)
    with pytest.raises(core.MixMapError, match="No usable plasmid/mass"):
        core.standardize_plate_csv(frame([[well, "Stock", 0]], wide), plate_type=6)


def test_parser_default_remains_unrestricted_for_internal_calculation():
    wells, long, _ = core.standardize_plate_csv(frame([["ZZ999", "Stock", 100]]))
    assert wells.Well.tolist() == ["ZZ999"]
    assert long.Mass_ng.tolist() == [100]


@pytest.mark.parametrize("plate_type,last_row,last_column", FORMATS)
@pytest.mark.parametrize("reagent", ["LT1", "L2000"])
def test_workflow_records_selected_geometry_and_preserves_calculations(plate_type, last_row, last_column, reagent):
    plate = b"Well,Plasmid,Mass (ng)\nA1,Stock,100\n"
    args = (plate, "selected.csv", STOCKS, "test", "test-time", [reagent], default_configs())
    result = generate(*args, plate_type=plate_type)
    reference = generate(*args)
    artifact = result["artifacts"][0]
    expected_rows = list("ABCDEFGH"[:ord(last_row) - ord("A") + 1])
    for metadata in (result, artifact):
        assert metadata["plate_type"] == plate_type
        assert metadata["plate_rows"] == expected_rows
        assert metadata["plate_columns"] == list(range(1, last_column + 1))
    pd.testing.assert_frame_equal(artifact["summary"], reference["artifacts"][0]["summary"])
    pd.testing.assert_frame_equal(artifact["details"], reference["artifacts"][0]["details"])
    assert artifact["bulk"] == reference["artifacts"][0]["bulk"]
    workbook = load_workbook(BytesIO(artifact["bytes"]))
    try:
        config = dict(workbook["Run config"].iter_rows(min_row=2, values_only=True))
        assert config["Plate type"] == get_plate_format(plate_type).label
    finally:
        workbook.close()


def test_workflow_default_is_48_and_single_plate_bounds_error_includes_filename():
    with pytest.raises(core.MixMapError, match="outside the selected plate type") as caught:
        generate(b"Well,Plasmid,Mass (ng)\nH12,Stock,100\n", "large.csv", STOCKS,
                 "test", "test-time", ["LT1"], default_configs())
    assert "large.csv: CSV row 2, well H12" in caught.value.details[0]
    assert "A1–F8" in caught.value.details[0]


def test_batch_collects_geometry_errors_before_any_workbook_is_generated(monkeypatch):
    def unexpected_write(*args, **kwargs):
        raise AssertionError("No workbook should be generated when any plate has out-of-bounds DNA")

    monkeypatch.setattr(core, "write_mix_map_workbook", unexpected_write)
    plates = [{"name": "good.csv", "bytes": b"Well,Plasmid,Mass (ng)\nA1,Stock,100\n"},
              {"name": "bad.csv", "bytes": b"Well,Plasmid,Mass (ng)\nC8,Stock,100\n"}]
    with pytest.raises(core.MixMapError, match="no output was made") as caught:
        generate_batch(plates, STOCKS, "test", "test-time", ["LT1"], default_configs(), plate_type=6)
    assert any("bad.csv (plate 2): CSV row 2, well C8" in detail for detail in caught.value.details)
    assert all("good.csv" not in detail for detail in caught.value.details)


@pytest.mark.parametrize("plate_type", [0, 384, "48", None, True])
def test_workflow_rejects_unsupported_plate_types(plate_type):
    with pytest.raises(core.MixMapError, match="Choose a supported plate type"):
        generate(b"Well,Plasmid,Mass (ng)\nA1,Stock,100\n", "plate.csv", STOCKS,
                 "test", "test-time", ["LT1"], default_configs(), plate_type=plate_type)
