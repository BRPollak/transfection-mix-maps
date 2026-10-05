"""Workbook regressions for literal labels and preparation/delivery volumes."""
import csv
from io import BytesIO, StringIO
import re

from openpyxl import load_workbook
import pytest

from workflow import default_configs, generate, generate_batch


def csv_bytes(headers, rows):
    stream = StringIO()
    writer = csv.writer(stream)
    writer.writerow(headers)
    writer.writerows(rows)
    return stream.getvalue().encode()


def table_cells(sheet):
    rows = sheet.iter_rows()
    headers = [cell.value for cell in next(rows)]
    return [dict(zip(headers, row)) for row in rows]


def texts(sheet):
    return [cell.value for row in sheet for cell in row if isinstance(cell.value, str)]


def volumes(text):
    return [float(value) for value in re.findall(r"(\d+(?:\.\d+)?)\s*[uµμ]L\b", text)]


@pytest.mark.parametrize("reagent", ["LT1", "L2000"])
def test_workbooks_keep_formula_like_and_numeric_identifiers_as_literal_text(tmp_path, reagent):
    names = ["=1+1", "=SUM(1,2)", "+1", "-1", "@SUM(1)", "00123", "NA", "None", "null", "nan"]
    rows = [(f"{'A' if index < 8 else 'B'}{index % 8 + 1}", name, 100)
            for index, name in enumerate(names)]
    plate = {"name": "=1+3.csv", "path": "=1+4", "bytes": csv_bytes(["Well", "Plasmid", "Mass (ng)"], rows)}
    sheet_label, source_label, timestamp = "=1+2", "=1+5", "=1+6"
    tables = {sheet_label: [["Plasmid", "Concentration"]] + [[name, 100] for name in names]}
    result = generate_batch([plate], tables, source_label, timestamp, [reagent], default_configs())
    artifact = result["artifacts"][0]
    workbook = load_workbook(BytesIO(artifact["bytes"]), data_only=False)
    try:
        assert not [(sheet.title, cell.coordinate) for sheet in workbook for row in sheet for cell in row
                    if cell.data_type == "f"]
        details = table_cells(workbook["Per-well details"])
        assert [row["Matched plasmid"].value for row in details] == names
        for row in details:
            assert row["Matched plasmid"].data_type == "s"
            assert row["Concentration source worksheet"].value == sheet_label
            assert row["Concentration source worksheet"].data_type == "s"
            assert row["Mass_ng"].value == 100
            assert row["Mass_ng"].data_type == "n"
        stocks = table_cells(workbook["Concentrations used"])
        assert {row["Plasmid"].value for row in stocks} == set(names)
        assert all(row["Plasmid"].data_type == "s" for row in stocks)
        assert all(row["Source worksheet"].value == sheet_label for row in stocks)
        config = {row["Setting"].value: row["Value"] for row in table_cells(workbook["Run config"])}
        for setting, expected in {
            "CSV file": plate["name"], "Plate CSV path": plate["path"],
            "Concentration source": source_label, "Concentrations loaded/refreshed at": timestamp,
        }.items():
            assert config[setting].value == expected
            assert config[setting].data_type == "s"
        title = workbook["Mix Map"]["A1"]
        assert title.value.startswith("=1+3")
        assert title.data_type == "s"
        map_text = "\n".join(texts(workbook["Mix Map"]))
        assert all(name in map_text for name in names)
    finally:
        workbook.close()


@pytest.mark.parametrize("reagent, prepared, displayed", [("LT1", 32.5, 32.6), ("L2000", 30, 30)])
def test_default_workbook_instructions_distinguish_prepared_and_delivered_volumes(tmp_path, reagent, prepared, displayed):
    result = generate(b"Well,Plasmid,Mass (ng)\nA1,Stock,100\n", "plate.csv",
                      {"Stocks": [["Plasmid", "Concentration"], ["Stock", 100]]},
                      "test", "test-time", [reagent], default_configs())
    artifact = result["artifacts"][0]
    workbook = load_workbook(BytesIO(artifact["bytes"]))
    try:
        notes = texts(workbook["Mix Map"])
        preparation_notes = [text for text in notes if text.lower().startswith("prepare ")]
        delivery_notes = [text for text in notes if text.lower().startswith("deliver ")]
        assert any(displayed in volumes(text) for text in preparation_notes)
        assert any(volumes(text)[0] == 25 for text in delivery_notes if volumes(text))
        row = table_cells(workbook["Well summary"])[0]
        dna_target = row["DNA mix target_uL"].value
        transfectant_target = row["Transfection mix target_uL"].value
        assert dna_target + transfectant_target == pytest.approx(prepared)
        assert (row["Total working DNA_uL"].value + row["DNA diluent_uL"].value
                + row["Reagent_uL"].value + row["Transfection diluent_uL"].value) == pytest.approx(prepared)
        if reagent == "L2000":
            assert dna_target == transfectant_target == 15
            bulk_notes = texts(workbook["Bulk transfectant mixes"])
            assert any("prepare" in text.lower() and "deliver" in text.lower()
                       and volumes(text)[:2] == [30, 25] for text in bulk_notes)
            assert all(mix["Per-well transfection mix_uL"] == 15 for mix in artifact["bulk"]["mixes"])
        else:
            assert dna_target == 32.5
            assert transfectant_target == 0
            assert artifact["bulk"] == {}
    finally:
        workbook.close()


@pytest.mark.parametrize("reagent", ["LT1", "L2000"])
def test_no_dna_wells_have_no_preparation_or_bulk_aliquot(tmp_path, reagent):
    plate = b"Well,Plasmid1,Mass1 (ng)\nA1,Stock,100\nA2,Stock,0\nA3,,\n"
    result = generate(plate, "plate.csv", {"Stocks": [["Plasmid", "Concentration"], ["Stock", 100]]},
                      "test", "test-time", [reagent], default_configs())
    artifact = result["artifacts"][0]
    workbook = load_workbook(BytesIO(artifact["bytes"]))
    try:
        summary = {row["Well"].value: row for row in table_cells(workbook["Well summary"])}
        volume_fields = [field for field in summary["A1"] if field.endswith("_uL")]
        for well in ["A2", "A3"]:
            assert summary[well]["Total DNA_ng"].value == 0
            assert all(summary[well][field].value == 0 for field in volume_fields)
            assert artifact["summary"].set_index("Well").loc[well, "Mix"] == ""
            assert not any(text.split()[:1] == [well] for text in texts(workbook["Mix Map"]))
            if reagent == "L2000":
                assert summary[well]["Bulk transfectant mix"].value in (None, "")
                assert summary[well]["Bulk mix symbol"].value in (None, "")
                assert all(well not in mix["Wells"] for mix in artifact["bulk"]["mixes"])
        if reagent == "L2000":
            assert sum(mix["Nonempty wells"] for mix in artifact["bulk"]["mixes"]) == 1
            assert sum(mix["Bulk total_uL"] for mix in artifact["bulk"]["mixes"]) == pytest.approx(18)
    finally:
        workbook.close()


def test_map_preparation_includes_custom_overages_without_changing_the_aliquot(tmp_path):
    configs = default_configs()
    configs["L2000"].update(final_volume_ul=40, dna_to_reagent_ratio_ul_per_ug=3,
                            well_overage_factor=1.5, bulk_overage_factor=1.7,
                            reagent_label="Custom reagent", diluent_label="Custom diluent")
    plate = csv_bytes(["Well", "Plasmid", "Mass (ng)"],
                      [("A1", "Stock", 100), ("A2", "Stock", 100), ("B1", "Stock", 300)])
    result = generate(plate, "plate.csv", {"Stocks": [["Plasmid", "Concentration"], ["Stock", 100]]},
                      "test", "test-time", ["L2000"], configs)
    workbook = load_workbook(BytesIO(result["artifacts"][0]["bytes"]))
    try:
        assert ("Bulk transfectant mix 1 ▲ | 100 ng DNA/well | 2 wells | "
                "Add 1.530 uL Custom reagent to 100.4 uL Custom diluent. | "
                "Add 30.0 uL to each assigned DNA mix.") in texts(workbook["Mix Map"])
        assert any(text.splitlines()[-2:] == ["Add 30.0 uL", "▲"] for text in texts(workbook["Mix Map"]))
    finally:
        workbook.close()
