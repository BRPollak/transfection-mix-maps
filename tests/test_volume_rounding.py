"""Pipette rounding is confined to printable recipes and map diagnostics."""
from io import BytesIO

from openpyxl import load_workbook
import pytest

import core
from workflow import default_configs, generate


@pytest.mark.parametrize("value, expected", [
    (None, "0.000"), (-1e-12, "0.000"), (-0.0, "0.000"), (0, "0.000"),
    (0.0004, "0.000"), (0.001, "0.002"), (0.0029, "0.002"), (0.003, "0.004"),
    (1.001, "1.002"), (1.005, "1.006"),
    (2.4989, "2.498"), (2.499, "2.500"), (2.5, "2.50"), (2.501, "2.50"),
    (2.505, "2.51"), (9.9949, "9.99"), (9.995, "10.00"),
    (10, "10.00"), (10.009, "10.00"), (10.01, "10.02"), (10.03, "10.04"),
    (19.99, "20.00"), (20, "20.0"), (20.0999, "20.0"), (20.1, "20.2"),
    (32.5, "32.6"), (199.9, "200.0"), (200, "200.0"), (200.0001, "200"),
    (200.5, "201"), (201.5, "202"),
])
def test_printable_volume_rounding_boundaries_and_ties(value, expected):
    assert core.format_volume_ul(value) == expected + " uL"


def text_cells(sheet):
    return [cell.value for row in sheet for cell in row if isinstance(cell.value, str)]


def table_records(sheet, header_row=1):
    rows = sheet.iter_rows(min_row=header_row, values_only=True)
    headers = next(rows)
    return [dict(zip(headers, row)) for row in rows]


@pytest.mark.parametrize("reagent", ["LT1", "L2000"])
def test_print_rounding_preserves_numeric_outputs_and_calculations(reagent):
    configs = default_configs()
    configs[reagent].update(final_volume_ul=401.1, well_overage_factor=1.0, bulk_overage_factor=1.0)
    result = generate(
        b"Well,Plasmid,Mass (ng)\nA1,Stock,100.1\n", "plate.csv",
        {"Stocks": [["Plasmid", "Concentration"], ["Stock", 100]]},
        "source", "timestamp", [reagent], configs,
    )
    artifact = result["artifacts"][0]
    workbook = load_workbook(BytesIO(artifact["bytes"]))
    try:
        well_text = next(text for text in text_cells(workbook["Mix Map"]) if text.startswith("A1 "))
        assert "1.002 uL  Stock" in well_text
        assert any(text.startswith("Deliver 401 uL") for text in text_cells(workbook["Mix Map"]))
        details = table_records(workbook["Per-well details"])[0]
        assert details["Working DNA volume_uL"] == pytest.approx(1.001)
        assert artifact["details"].iloc[0]["Working DNA volume_uL"] == pytest.approx(1.001)
        for row in workbook["Per-well details"].iter_rows(min_row=2):
            assert row[5].number_format == row[6].number_format == "0.000"
        numeric_summary = table_records(workbook["Well summary"])[0]
        for field, value in numeric_summary.items():
            if field.endswith("_uL"):
                assert value == pytest.approx(artifact["summary"].iloc[0][field], rel=1e-14)
        assert dict(workbook["Run config"].iter_rows(min_row=2, values_only=True))["Final volume_uL"] == 401.1
        if reagent == "L2000":
            assert "Add 201 uL" in well_text
            bulk_sheet = workbook["Bulk transfectant mixes"]
            assert any("201 uL DNA mix" in text and "201 uL of this bulk" in text
                       for text in text_cells(bulk_sheet))
            assert any("0.200 uL Lipofectamine 2000 + 200 uL Diluent (Opti-MEM) = 201 uL total"
                       in text for text in text_cells(bulk_sheet))
            header_row = next(cell.row for cell in bulk_sheet["A"] if cell.value == "Mix")
            recipe = table_records(bulk_sheet, header_row)[0]
            for field, value in recipe.items():
                if field and field.endswith("_uL"):
                    assert value == pytest.approx(artifact["bulk"]["mixes"][0][field], rel=1e-14)
    finally:
        workbook.close()


@pytest.mark.parametrize("reagent", ["LT1", "L2000"])
def test_map_warning_rounds_original_values_without_rewriting_labels_or_diagnostics(reagent):
    configs = default_configs()
    configs[reagent]["well_overage_factor"] = 1.0
    name = "Dose 0.003 uL"
    result = generate(
        f"Well,Plasmid,Mass (ng)\nA1,{name},0.29\n".encode(), "plate.csv",
        {"Stocks": [["Plasmid", "Concentration"], [name, 100]]},
        "source", "timestamp", [reagent], configs,
    )
    workbook = load_workbook(BytesIO(result["artifacts"][0]["bytes"]))
    try:
        raw_warning = next(warning for warning in result["warnings"] if "DNA component volume(s)" in warning)
        assert f"A1 {name} 0.003 uL" in raw_warning
        map_warning = next(text for text in text_cells(workbook["Mix Map"]) if "DNA component volume(s)" in text)
        # Rounding the old 0.003 diagnostic would yield 0.004; the source 0.0029
        # must round directly to 0.002, without changing the number in the name.
        assert f"A1 {name} 0.002 uL" in map_warning
        assert "below 0.200 uL" in map_warning
        assert raw_warning in text_cells(workbook["Warnings"])
        assert raw_warning in text_cells(workbook["Run config"])
        well_text = next(text for text in text_cells(workbook["Mix Map"]) if text.startswith("A1 "))
        assert f"0.002 uL  {name}" in well_text
    finally:
        workbook.close()
