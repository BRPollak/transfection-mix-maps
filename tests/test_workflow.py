import csv
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

import core
import workflow
from sources import read_plate
from workflow import default_configs, generate, generate_batch, save_artifact

FIXTURES = Path(__file__).resolve().parent / "fixtures"


def inputs():
    """Synthetic inputs exist only as test fixtures, never as app input options."""
    plate = (FIXTURES / "plate.csv").read_bytes()
    with (FIXTURES / "concentrations.csv").open(newline="", encoding="utf-8-sig") as handle:
        tables = {"concentrations": list(csv.reader(handle))}
    return plate, tables


def calculate(reagent, raw=None):
    plate, tables = inputs()
    wells, long, _ = core.standardize_plate_csv(read_plate(plate) if raw is None else raw)
    _, lookup, _ = core.load_plasmid_concentrations(tables, used_plasmids=long.Plasmid.tolist())
    matched, _ = core.match_concentrations(long, lookup)
    return core.calculate_mix(wells, matched, default_configs()[reagent])


def test_lt1_known_volumes():
    details, summary, bulk, warnings = calculate("LT1")
    a1 = summary.set_index("Well").loc["A1"]
    assert a1["Total DNA_ng"] == 150
    assert a1["Total working DNA_uL"] == pytest.approx(2.6)
    assert a1["Reagent_uL"] == pytest.approx(0.6825)
    assert a1["DNA diluent_uL"] == pytest.approx(29.2175)
    assert a1["DNA mix target_uL"] == pytest.approx(32.5)
    assert bulk == {}


def test_l2000_known_volumes():
    assert default_configs()["L2000"]["final_volume_ul"] == 25
    _, summary, bulk, _ = calculate("L2000")
    a1 = summary.set_index("Well").loc["A1"]
    assert a1["DNA mix target_uL"] == pytest.approx(15)
    assert a1["Total working DNA_uL"] == pytest.approx(2.4)
    assert a1["DNA diluent_uL"] == pytest.approx(12.6)
    assert a1["Reagent_uL"] == pytest.approx(0.36)
    assert a1["Transfection diluent_uL"] == pytest.approx(14.64)
    assert len(bulk["mixes"]) == 1
    mix = bulk["mixes"][0]
    assert mix["Mix"] == "Bulk transfectant mix 1"
    assert mix["Wells"] == ["A1", "A2", "B1", "B2"]
    assert mix["Bulk reagent_uL"] == pytest.approx(1.728)
    assert mix["Bulk total_uL"] == pytest.approx(72)


def test_long_and_wide_match():
    long = pd.DataFrame([
        ["A1", "Example_A", 100], ["A1", "Example_B", 50],
        ["A2", "Example_A", 100], ["A2", "Example_C", 50],
        ["B1", "Example_B", 150], ["B2", "Example_C", 150],
    ], columns=["Well", "Plasmid", "Mass (ng)"])
    pd.testing.assert_frame_equal(calculate("LT1")[1], calculate("LT1", long)[1])


def test_generation_ignores_unrelated_sheet_format_problems(tmp_path):
    plate = b"Well,Plasmid,Mass (ng)\nA1,Used,100\n"
    tables = {
        "Stocks": [["Plasmid", "Concentration"], ["Used", ""], ["Used", "50"], ["Unused", ""],
                   ["Other", "text100"], ["", "10"], ["Conflict", "20"], ["Conflict", "30"]],
        "Partial": [["Plasmid", "Location"], ["Other", "Freezer"]],
        "Different units": [["Plasmid", "Concentration (ug/uL)"], ["Unused", "5"]],
        "Notes": [["Date", "Operator"], ["Oct 4", "Lab"]],
        "Empty": [],
    }
    result = generate(plate, "plate.csv", tables, "test", "test", ["LT1"], default_configs())
    assert len(result["artifacts"]) == 1
    assert not result["warnings"]
    workbook = load_workbook(BytesIO(result["artifacts"][0]["bytes"]))
    stocks = list(workbook["Concentrations used"].iter_rows(min_row=2, values_only=True))
    assert len(stocks) == 1
    assert stocks[0][0] == "Used"


def test_generation_reports_all_unusable_plate_plasmids_and_makes_no_output(tmp_path):
    plate = b"Well,Plasmid,Mass (ng)\nA1,Absent,100\nA2,Blank,100\nA3,Invalid,100\nA4,Conflict,100\n"
    tables = {"Stocks": [["Plasmid", "Concentration"], ["Blank", ""], ["Invalid", "text100"],
                          ["Conflict", "50"], ["Conflict", "60"]]}
    with pytest.raises(core.MixMapError, match="no output was made") as caught:
        generate(plate, "plate.csv", tables, "test", "test", ["LT1"], default_configs())
    assert len(caught.value.details) == 4
    assert any("Absent: not found" in detail for detail in caught.value.details)
    assert any("Blank: concentration is missing" in detail for detail in caught.value.details)
    assert any("Invalid: concentration must be" in detail for detail in caught.value.details)
    assert any("Conflict: conflicting concentrations" in detail for detail in caught.value.details)
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("tables, reason", [
    ({"Stocks": [["Plasmid"], ["A"]]}, "concentration column is missing"),
    ({"Stocks": [["Plasmid", "Concentration (ug/uL)"], ["A", "50"]]}, "units must be ng/µL"),
    ({"Stocks": [["Plasmid", "Concentration", "Concentration"], ["A", "50", "60"]]}, "column is ambiguous"),
])
def test_used_plasmid_requires_a_usable_concentration_column(tmp_path, tables, reason):
    plate = b"Well,Plasmid,Mass (ng)\nA1,A,100\n"
    with pytest.raises(core.MixMapError) as caught:
        generate(plate, "plate.csv", tables, "test", "test", ["LT1"], default_configs())
    assert len(caught.value.details) == 1
    assert caught.value.details[0].startswith("A:")
    assert reason in caught.value.details[0]
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("bad", ["1e999", "0"])
def test_used_plasmid_rejects_nonfinite_or_nonpositive_concentration(bad):
    tables = {"Stocks": [["Plasmid", "Concentration"], ["A", bad]]}
    with pytest.raises(core.MixMapError) as caught:
        core.load_plasmid_concentrations(tables, used_plasmids=["A"])
    assert "A: concentration must be a finite, positive number" in caught.value.details[0]


@pytest.mark.parametrize("concentration", ["50 ng per uL", "5e1"])
def test_concentration_parser_supports_numbers_and_explicit_ng_per_ul(concentration):
    tables = {"Stocks": [["Plasmid", "Concentration"], ["A", concentration]]}
    _, lookup, _ = core.load_plasmid_concentrations(tables, used_plasmids=["A"])
    assert lookup["A"]["Concentration_ng_per_uL"] == 50


def test_case_matching_preserves_exact_names_and_rejects_ambiguous_fallback():
    tables = {"Stocks": [["Plasmid", "Concentration"], ["Stock", "50"], ["stock", "100"]]}
    _, lookup, _ = core.load_plasmid_concentrations(tables, used_plasmids=["Stock"])
    assert lookup["Stock"]["Concentration_ng_per_uL"] == 50
    with pytest.raises(core.MixMapError) as caught:
        core.load_plasmid_concentrations(tables, used_plasmids=["STOCK"])
    assert "STOCK: ambiguous name match" in caught.value.details[0]


def test_case_insensitive_matching_keeps_stock_provenance():
    tables = {"Stocks": [["Plasmid", "Concentration"], ["Stock", "50"]]}
    _, lookup, _ = core.load_plasmid_concentrations(tables, used_plasmids=["STOCK"])
    _, long, _ = core.standardize_plate_csv(pd.DataFrame(
        [["A1", "STOCK", 100]], columns=["Well", "Plasmid", "Mass (ng)"]))
    matched, warnings = core.match_concentrations(long, lookup)
    assert matched.iloc[0]["Matched plasmid"] == "Stock"
    assert matched.iloc[0]["Concentration source worksheet"] == "Stocks"
    assert warnings


def test_duplicate_wide_wells_stop():
    plate, _ = inputs()
    raw = read_plate(plate)
    with pytest.raises(core.MixMapError, match="Duplicate well"):
        core.standardize_plate_csv(pd.concat([raw, raw.iloc[:1]], ignore_index=True))


def test_nonfinite_dna_mass_does_not_publish_output(tmp_path):
    plate = b"Well,Plasmid,Mass (ng)\nA1,A,1e999\n"
    tables = {"Stocks": [["Plasmid", "Concentration"], ["A", "50"]]}
    with pytest.raises(core.MixMapError, match="unreadable DNA mass"):
        generate(plate, "plate.csv", tables, "test", "test-time", ["LT1"], default_configs())
    assert not list(tmp_path.iterdir())


def test_impossible_volume_does_not_publish_partial_run(tmp_path):
    plate, tables = inputs()
    configs = default_configs()
    configs["LT1"]["final_volume_ul"] = 0.1
    with pytest.raises(core.MixMapError, match="more volume"):
        generate(plate, "plate.csv", tables, "test", "test-time", ["LT1"], configs)
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("reagent", ["LT1", "L2000"])
def test_workbooks_and_repeated_runs_stay_in_memory(tmp_path, monkeypatch, reagent):
    plate, tables = inputs()
    monkeypatch.chdir(tmp_path)

    def unexpected_filesystem_output(*args, **kwargs):
        raise AssertionError("Generation must not stage or publish filesystem workbooks")

    monkeypatch.setattr(workflow.tempfile, "mkstemp", unexpected_filesystem_output)
    monkeypatch.setattr(workflow.os, "link", unexpected_filesystem_output)
    monkeypatch.setattr(core.os, "makedirs", unexpected_filesystem_output)
    args = (plate, "plate.csv", tables, "test source", "2026-10-04T12:00:00-07:00", [reagent], default_configs())
    first = generate(*args)
    first_bytes = first["artifacts"][0]["bytes"]
    second = generate(*args)
    assert "folder" not in first
    assert "folder" not in second
    assert first["artifacts"][0]["name"] != second["artifacts"][0]["name"]
    assert first["artifacts"][0]["bytes"] == first_bytes
    assert len(first["artifacts"]) == 1
    assert first["artifacts"][0]["reagent"] == reagent
    assert not list(tmp_path.iterdir())
    for artifact in first["artifacts"]:
        assert "path" not in artifact
        wb = load_workbook(BytesIO(artifact["bytes"]))
        assert {"Mix Map", "Per-well details", "Concentrations used", "Well summary", "Run config"} <= set(wb.sheetnames)
        assert wb["Mix Map"].page_setup.orientation == "landscape"
        assert str(wb["Mix Map"].page_setup.paperSize) == wb["Mix Map"].PAPERSIZE_LETTER
        assert wb["Mix Map"].page_setup.fitToWidth == 1
        config = dict(wb["Run config"].iter_rows(min_row=2, values_only=True))
        assert config["Concentrations loaded/refreshed at"] == args[4]
        assert config["Duplicate concentration policy"] == "error"
        assert len(config["Concentration snapshot SHA256"]) == 64
        wb.close()


def test_download_saves_only_selected_workbook_and_exact_generated_bytes(tmp_path):
    plate, tables = inputs()
    result = generate_batch([{"name": name, "bytes": plate} for name in ["one.csv", "two.csv"]],
                            tables, "test", "test", ["LT1"], default_configs())
    selected = tmp_path / "selected"
    downloads = tmp_path / "Downloads"
    selected.mkdir()
    downloads.mkdir()
    artifact = result["artifacts"][1]
    initial_keys = set(artifact)
    saved = save_artifact(artifact, selected)
    destination = Path(saved["path"])
    assert destination == selected / artifact["name"]
    assert destination.read_bytes() == artifact["bytes"]
    assert saved["warnings"] == []
    assert set(artifact) == initial_keys
    assert "path" not in artifact
    assert list(selected.iterdir()) == [destination]
    assert not list(downloads.iterdir())
    workbook = load_workbook(destination)
    assert "Mix Map" in workbook.sheetnames
    workbook.close()


def test_missing_output_folder_is_rejected_only_on_download_and_not_created(tmp_path):
    plate, tables = inputs()
    artifact = generate(plate, "plate.csv", tables, "test", "test", ["LT1"], default_configs())["artifacts"][0]
    missing = tmp_path / "missing" / "nested"
    with pytest.raises(core.MixMapError, match="existing output folder"):
        save_artifact(artifact, missing)
    assert artifact["bytes"]
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("selection", [None, "", "   "])
def test_download_requires_an_explicit_folder(tmp_path, monkeypatch, selection):
    monkeypatch.chdir(tmp_path)
    with pytest.raises(core.MixMapError, match="existing output folder"):
        save_artifact({"name": "map.xlsx", "bytes": b"workbook"}, selection)
    assert not list(tmp_path.iterdir())


def test_download_revalidates_folder_after_generation(tmp_path):
    plate, tables = inputs()
    selected = tmp_path / "selected"
    selected.mkdir()
    artifact = generate(plate, "plate.csv", tables, "test", "test", ["LT1"], default_configs())["artifacts"][0]
    selected.rmdir()
    with pytest.raises(core.MixMapError, match="existing output folder"):
        save_artifact(artifact, selected)
    assert not selected.exists()


def test_repeated_downloads_preserve_previous_file(tmp_path):
    artifact = {"name": "map.xlsx", "bytes": b"generated bytes"}
    existing = tmp_path / artifact["name"]
    existing.write_bytes(b"prior workbook")
    first = save_artifact(artifact, tmp_path)
    second = save_artifact(artifact, tmp_path)
    assert existing.read_bytes() == b"prior workbook"
    assert first["path"] != second["path"]
    assert Path(first["path"]).read_bytes() == Path(second["path"]).read_bytes() == artifact["bytes"]
    assert len(list(tmp_path.iterdir())) == 3


def test_publication_collision_preserves_existing_file_and_retries(tmp_path, monkeypatch):
    artifact = {"name": "map.xlsx", "bytes": b"generated bytes"}
    real_link = workflow.os.link
    collisions = []

    def collide_once(source, destination):
        if not collisions:
            destination.write_bytes(b"another process saved this file")
            collisions.append(destination)
        return real_link(source, destination)

    monkeypatch.setattr(workflow.os, "link", collide_once)
    result = save_artifact(artifact, tmp_path)
    assert collisions[0].read_bytes() == b"another process saved this file"
    saved = Path(result["path"])
    assert saved != collisions[0]
    assert saved.read_bytes() == artifact["bytes"]
    assert len(list(tmp_path.iterdir())) == 2


@pytest.mark.parametrize("reagents", [["LT1", "L2000"], ["unknown"]])
def test_invalid_reagent_selection_does_not_publish_output(tmp_path, reagents):
    plate, tables = inputs()
    with pytest.raises(core.MixMapError, match="exactly one transfectant"):
        generate(plate, "plate.csv", tables, "test", "test", reagents,
                 default_configs())
    assert not list(tmp_path.iterdir())


def test_batch_creates_distinct_workbooks_for_same_named_plates(tmp_path):
    first = b"Well,Plasmid,Mass (ng)\nA1,A,100\nA1,B,50\n"
    second = b"Well,Plasmid,Mass (ng)\nB2,A,150\n"
    plates = [{"name": "plate.csv", "bytes": first, "path": "/first/plate.csv"},
              {"name": "plate.csv", "bytes": second, "path": "/second/plate.csv"}]
    tables = {"A stocks": [["Plasmid", "Concentration"], ["A", "50"]],
              "B stocks": [["Plasmid", "Concentration"], ["B", "100"]]}
    result = generate_batch(plates, tables, "test", "test-time", ["LT1"], default_configs())
    assert result["plate_count"] == 2
    assert result["wells"] == 2
    assert result["entries"] == 3
    assert result["plasmids"] == 2
    assert len(result["artifacts"]) == 2
    assert len({artifact["name"] for artifact in result["artifacts"]}) == 2
    assert [artifact["plate_path"] for artifact in result["artifacts"]] == [plate["path"] for plate in plates]
    assert all(artifact["plate_name"] == "plate.csv" for artifact in result["artifacts"])
    assert result["artifacts"][0]["summary"]["Well"].tolist() == ["A1"]
    assert result["artifacts"][1]["summary"]["Well"].tolist() == ["B2"]
    for artifact, plate in zip(result["artifacts"], plates):
        assert "path" not in artifact
        workbook = load_workbook(BytesIO(artifact["bytes"]))
        config = dict(workbook["Run config"].iter_rows(min_row=2, values_only=True))
        assert config["Plate CSV path"] == plate["path"]
        assert config["Plate CSV SHA256"] == workflow.hashlib.sha256(plate["bytes"]).hexdigest()
    first_workbook = load_workbook(BytesIO(result["artifacts"][0]["bytes"]))
    stocks = first_workbook["Concentrations used"].iter_rows(min_row=2, values_only=True)
    assert [(row[0], row[2], row[3]) for row in stocks] == [("A", "A stocks", 2), ("B", "B stocks", 2)]
    assert not list(tmp_path.iterdir())


def test_batch_accepts_five_plates(tmp_path):
    count = 5
    plate, tables = inputs()
    result = generate_batch([{"name": f"plate-{index}.csv", "bytes": plate} for index in range(count)],
                            tables, "test", "test-time", ["LT1"], default_configs())
    assert result["plate_count"] == count
    assert len(result["artifacts"]) == count
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("count", [0, 6])
def test_batch_size_limits_are_enforced_before_output(tmp_path, count):
    plate, tables = inputs()
    with pytest.raises(core.MixMapError, match="between 1 and 5"):
        generate_batch([{"name": f"plate-{index}.csv", "bytes": plate} for index in range(count)],
                       tables, "test", "test-time", ["LT1"], default_configs())
    assert not list(tmp_path.iterdir())


def test_batch_collects_plate_validation_errors_without_writing_any_workbooks(tmp_path, monkeypatch):
    plates = [
        {"name": "good.csv", "bytes": b"Well,Plasmid,Mass (ng)\nA1,A,100\n"},
        {"name": "missing.csv", "bytes": b"Well,Plasmid,Mass (ng)\nA1,Absent,100\nA2,Blank,100\n"},
        {"name": "malformed.csv", "bytes": b"Well,Plasmid,Mass (ng)\nNotAWell,A,100\n"},
    ]
    tables = {"Stocks": [["Plasmid", "Concentration"], ["A", "50"], ["Blank", ""]]}

    def unexpected_write(*args):
        raise AssertionError("No workbook may be staged when any plate has validation errors")

    monkeypatch.setattr(core, "write_mix_map_workbook", unexpected_write)
    with pytest.raises(core.MixMapError, match="no output was made") as caught:
        generate_batch(plates, tables, "test", "test-time", ["LT1"], default_configs())
    details = "\n".join(caught.value.details)
    assert "missing.csv (plate 2): Absent: not found" in details
    assert "missing.csv (plate 2): Blank: concentration is missing" in details
    assert "malformed.csv (plate 3): Malformed well" in details
    assert "good.csv" not in details
    assert not list(tmp_path.iterdir())


def test_batch_generation_failure_never_writes_a_workbook(tmp_path, monkeypatch):
    plate, tables = inputs()
    existing = tmp_path / "previous.xlsx"
    existing.write_bytes(b"prior output")
    real_write = core.write_mix_map_workbook
    writes = []

    def fail_second_write(stream, *args):
        assert isinstance(stream, BytesIO)
        writes.append(stream)
        if len(writes) == 2:
            stream.write(b"partial")
            raise OSError("second workbook could not be generated")
        return real_write(stream, *args)

    monkeypatch.setattr(core, "write_mix_map_workbook", fail_second_write)
    with pytest.raises(OSError, match="second workbook"):
        generate_batch([{"name": name, "bytes": plate} for name in ["one.csv", "two.csv"]],
                       tables, "test", "test-time", ["LT1"], default_configs())
    assert len(writes) == 2
    assert existing.read_bytes() == b"prior output"
    assert list(tmp_path.iterdir()) == [existing]


def test_batch_warnings_identify_their_plate(tmp_path):
    plates = [{"name": "small.csv", "bytes": b"Well,Plasmid,Mass (ng)\nA1,A,1\n"},
              {"name": "normal.csv", "bytes": b"Well,Plasmid,Mass (ng)\nA1,A,100\n"}]
    tables = {"Stocks": [["Plasmid", "Concentration"], ["A", "100"]]}
    result = generate_batch(plates, tables, "test", "test-time", ["LT1"], default_configs())
    assert result["warnings"]
    assert all(warning.startswith("small.csv (plate 1): ") for warning in result["warnings"])
    assert "plate_path" not in result["artifacts"][0]
    workbook = load_workbook(BytesIO(result["artifacts"][0]["bytes"]))
    assert "Warnings" in workbook.sheetnames


def test_successful_download_reports_temporary_cleanup_failure_as_warning(tmp_path, monkeypatch):
    artifact = {"name": "map.xlsx", "bytes": b"generated bytes"}
    real_unlink = Path.unlink
    cleanup_attempts = []

    def fail_temporary_cleanup(path, *args, **kwargs):
        if path.parent == tmp_path and path.name.startswith(".mixmap-"):
            cleanup_attempts.append(path)
            raise PermissionError("temporary file is locked")
        return real_unlink(path, *args, **kwargs)

    monkeypatch.setattr(Path, "unlink", fail_temporary_cleanup)
    result = save_artifact(artifact, tmp_path)
    assert Path(result["path"]).read_bytes() == artifact["bytes"]
    assert len(cleanup_attempts) == 1
    assert cleanup_attempts[0].exists()
    assert any("Workbook was saved" in warning and str(cleanup_attempts[0]) in warning
               for warning in result["warnings"])


@pytest.mark.parametrize("failure_stage", ["write", "publish"])
def test_failed_download_cleans_up_and_preserves_existing_files(tmp_path, monkeypatch, failure_stage):
    artifact = {"name": "map.xlsx", "bytes": b"generated bytes"}
    existing = tmp_path / "previous.xlsx"
    existing.write_bytes(b"prior output")

    def fail(*args, **kwargs):
        raise PermissionError("download failed")

    if failure_stage == "write":
        monkeypatch.setattr(workflow.os, "fsync", fail)
    else:
        monkeypatch.setattr(workflow.os, "link", fail)
    with pytest.raises(PermissionError, match="download failed"):
        save_artifact(artifact, tmp_path)
    assert existing.read_bytes() == b"prior output"
    assert list(tmp_path.iterdir()) == [existing]
    assert "path" not in artifact


def test_temporary_cleanup_failure_does_not_mask_primary_save_error(tmp_path, monkeypatch):
    artifact = {"name": "map.xlsx", "bytes": b"generated bytes"}
    real_unlink = Path.unlink
    primary_error = PermissionError("workbook could not be published")
    cleanup_attempts = []

    def fail_publication(*args, **kwargs):
        raise primary_error

    def fail_temporary_cleanup(path, *args, **kwargs):
        if path.parent == tmp_path and path.name.startswith(".mixmap-"):
            cleanup_attempts.append(path)
            raise PermissionError("temporary file is locked")
        return real_unlink(path, *args, **kwargs)

    monkeypatch.setattr(workflow.os, "link", fail_publication)
    monkeypatch.setattr(Path, "unlink", fail_temporary_cleanup)
    with pytest.raises(PermissionError, match="workbook could not be published") as caught:
        save_artifact(artifact, tmp_path)
    assert caught.value is primary_error
    assert len(cleanup_attempts) == 1
    assert len(caught.value.cleanup_warnings) == 1
    assert str(cleanup_attempts[0]) in caught.value.cleanup_warnings[0]
    assert not getattr(caught.value, "outputs_created", False)
    assert all(path.name.startswith(".mixmap-") for path in tmp_path.iterdir())


@pytest.mark.parametrize("name", ["../escape.xlsx", "/absolute.xlsx", "nested/map.xlsx", "nested\\map.xlsx", "", ".."])
def test_download_rejects_unsafe_names(tmp_path, name):
    with pytest.raises(core.MixMapError, match="workbook is unavailable"):
        save_artifact({"name": name, "bytes": b"generated bytes"}, tmp_path)
    assert not list(tmp_path.iterdir())
