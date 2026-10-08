"""Printable maps use the selected physical geometry and whole-plate recipes."""

import csv
from io import BytesIO, StringIO

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.utils import get_column_letter
import pytest

import core
from plate_formats import PLATE_FORMATS
from workflow import default_configs, generate


STOCKS = {"Stocks": [["Plasmid", "Concentration"], ["Stock", 100]]}


def workbook_for(rows, plate_type, reagent="L2000", stocks=None):
    stream = StringIO()
    writer = csv.writer(stream)
    writer.writerow(["Well", "Plasmid", "Mass (ng)"])
    writer.writerows(rows)
    result = generate(
        stream.getvalue().encode(), "plate.csv", stocks or STOCKS, "test", "test-time",
        [reagent], default_configs(), plate_type=plate_type,
    )
    artifact = result["artifacts"][0]
    return load_workbook(BytesIO(artifact["bytes"])), artifact


def grid_header(sheet):
    return sheet[sheet.freeze_panes].row - 1


def grid_positions(sheet):
    header_row = grid_header(sheet)
    return {
        f"{sheet.cell(row, 1).value}{sheet.cell(header_row, column).value}": sheet.cell(row, column)
        for row in range(header_row + 1, sheet.max_row + 1)
        for column in range(2, sheet.max_column + 1)
        if not isinstance(sheet.cell(row, column), MergedCell)
    }


def texts(sheet):
    return [cell.value for row in sheet for cell in row if isinstance(cell.value, str)]


def assert_grayscale_workbook(workbook):
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                colors = [cell.font.color]
                if cell.fill.patternType == "solid":
                    color = cell.fill.fgColor
                    assert color.type == "rgb"
                    assert int(color.rgb[-6:-4], 16) >= 0xD4
                    colors.append(color)
                colors.extend(side.color for side in (
                    cell.border.left, cell.border.right, cell.border.top, cell.border.bottom,
                ) if side is not None)
                for color in colors:
                    if color is not None and color.type == "rgb":
                        rgb = color.rgb[-6:]
                        assert rgb[:2] == rgb[2:4] == rgb[4:], (sheet.title, cell.coordinate, rgb)


@pytest.mark.parametrize("plate_type", [96, 48, 24, 12, 6, 1])
@pytest.mark.parametrize("reagent", ["LT1", "L2000"])
def test_sparse_workbooks_show_every_supported_position_exactly_once(plate_type, reagent):
    plate_format = PLATE_FORMATS[plate_type]
    active_wells = {"A1", f"{plate_format.rows[-1]}{plate_format.columns[-1]}"}
    workbook, artifact = workbook_for([(well, "Stock", 100) for well in sorted(active_wells)], plate_type, reagent)
    rich = load_workbook(BytesIO(artifact["bytes"]), rich_text=True)
    try:
        assert_grayscale_workbook(workbook)
        maps = [sheet for sheet in workbook if sheet.title.startswith("Mix Map")]
        assert [sheet.title for sheet in maps] == [name for name, _ in plate_format.map_sections]
        found = []
        for sheet, (_, expected_rows) in zip(maps, plate_format.map_sections):
            header_row = grid_header(sheet)
            assert sheet.cell(header_row, 1).value is None
            assert "Concentration source" not in sheet["A2"].value
            assert sheet["A2"].value.startswith(f"{plate_format.label} ({plate_format.well_range}) | Generated ")
            assert [cell.value for cell in sheet[header_row][1:]
                    if not isinstance(cell, MergedCell)] == list(plate_format.columns)
            assert [sheet.cell(row, 1).value for row in range(header_row + 1, sheet.max_row + 1)] == list(expected_rows)
            labels = [cell for cell in sheet[header_row] if not isinstance(cell, MergedCell)]
            labels.extend(sheet.cell(row, 1) for row in range(header_row + 1, sheet.max_row + 1))
            assert all(cell.fill.fgColor.rgb == "00D4D4D4" for cell in labels)
            positions = grid_positions(sheet)
            found.extend(positions)
            for well, cell in positions.items():
                assert cell.fill.fgColor.rgb == ("00F7F7F7" if well in active_wells else "00E6E6E6")
                assert cell.alignment.wrap_text
                if well in active_wells:
                    styled = rich[sheet.title][cell.coordinate]
                    assert isinstance(styled.value, CellRichText)
                    assert str(styled.value) == cell.value
                    label = styled.value[0]
                    assert isinstance(label, TextBlock)
                    assert label.text == well
                    assert label.font.b is True
                    assert not styled.font.b
                    assert not any(part.font.b for part in styled.value[1:] if isinstance(part, TextBlock))
            assert sheet.page_setup.orientation == "landscape"
            assert str(sheet.page_setup.paperSize) == sheet.PAPERSIZE_LETTER
            assert sheet.page_setup.fitToWidth == sheet.page_setup.fitToHeight == 1
            assert sheet.sheet_properties.pageSetUpPr.fitToPage
            expected_last_column = 9 if plate_type == 1 else len(plate_format.columns) + 1
            assert str(sheet.print_area).endswith(f"$A$1:${get_column_letter(expected_last_column)}${sheet.max_row}")
            assert sheet.freeze_panes == f"B{header_row + 1}"
        assert len(found) == len(set(found)) == plate_type
        assert set(found) == {f"{row}{column}" for row in plate_format.rows for column in plate_format.columns}
        assert {"Per-well details", "Concentrations used", "Well summary", "Run config"} <= set(workbook.sheetnames)
        assert ("Bulk transfectant mixes" in workbook.sheetnames) == (reagent == "L2000")
        if plate_type == 1:
            assert grid_positions(maps[0])["A1"].value.split()[0] == "A1"
        settings = dict(workbook["Run config"].iter_rows(min_row=2, max_col=2, values_only=True))
        assert settings["Plate type"] == plate_format.label
        assert settings["Plate wells"] == plate_type
        assert settings["Plate rows"] == ", ".join(plate_format.rows)
        assert settings["Plate columns"] == len(plate_format.columns)
        assert settings["Plate well range"] == plate_format.well_range
        if plate_type == 96:
            assert "A-D and E-H" in settings["Print layout"]
    finally:
        workbook.close()
        rich.close()


@pytest.mark.parametrize("plate_type", [24, 12, 6, 1])
def test_small_formats_retain_width_for_printed_preparation_notes(plate_type):
    workbook, _ = workbook_for([("A1", "Stock", 100)], plate_type)
    try:
        sheet = workbook["Mix Map"]
        assert sheet.max_column == (9 if plate_type == 1 else len(PLATE_FORMATS[plate_type].columns) + 1)
        assert sum(sheet.column_dimensions[get_column_letter(column)].width
                   for column in range(2, sheet.max_column + 1)) == pytest.approx(176)
    finally:
        workbook.close()


@pytest.mark.parametrize("reagent", ["LT1", "L2000"])
def test_dish_merges_printable_columns_into_one_labeled_well(reagent):
    workbook, _ = workbook_for([("A1", "Stock", 100)], 1, reagent)
    try:
        sheet = workbook["Mix Map"]
        header_row = grid_header(sheet)
        assert set(grid_positions(sheet)) == {"A1"}
        assert sheet.cell(header_row, 2).value == 1
        assert sheet.cell(header_row + 1, 2).value.startswith("A1 ")
        assert f"B{header_row}:I{header_row}" in sheet.merged_cells
        assert f"B{header_row + 1}:I{header_row + 1}" in sheet.merged_cells
        assert "A1:I1" in sheet.merged_cells
        assert all(sheet.column_dimensions[column].width == 22 for column in "BCDEFGHI")
        assert sheet.cell(header_row + 1, 9).border.right.style == "medium"
    finally:
        workbook.close()


def test_96_well_halves_repeat_whole_plate_recipes_and_share_symbols():
    rows = [(f"{row}{column}", "Stock", column * 100) for row in ("A", "H") for column in range(1, 6)]
    workbook, artifact = workbook_for(rows, 96)
    try:
        top, bottom = workbook["Mix Map (A-D)"], workbook["Mix Map (E-H)"]
        assert "Rows A-D" in top["A1"].value
        assert "Rows E-H" in bottom["A1"].value
        recipes = artifact["bulk"]["mixes"]
        assert len(recipes) == 5
        legends = []
        for sheet in (top, bottom):
            assert any("Whole-plate totals" in text and "once" in text for text in texts(sheet))
            legends.append([text for text in texts(sheet) if text.startswith("Bulk transfectant mix ")])
        assert legends[0] == legends[1]
        for recipe in recipes:
            assert recipe["Nonempty wells"] == 2
            assert recipe["Bulk total_uL"] == pytest.approx(36)
            assert any(f"Add {core.format_volume_ul(recipe['Bulk reagent_uL'])} Lipofectamine 2000"
                       in legend for legend in legends[0])
            for well in recipe["Wells"]:
                cell = grid_positions(top if well.startswith("A") else bottom)[well]
                assert cell.value.splitlines()[-1] == recipe["Symbol"]
                assert "Add 15.00 uL" in cell.value
    finally:
        workbook.close()


def test_96_well_empty_bottom_half_still_has_the_same_recipes_and_notes():
    workbook, _ = workbook_for([("A1", "Stock", 100)], 96)
    try:
        top, bottom = workbook["Mix Map (A-D)"], workbook["Mix Map (E-H)"]
        assert all(cell.value is None for cell in grid_positions(bottom).values())
        assert [top.cell(row, 1).value for row in range(2, grid_header(top))] == [
            bottom.cell(row, 1).value for row in range(2, grid_header(bottom))
        ]
    finally:
        workbook.close()


def test_dense_96_well_long_labels_reserve_height_for_wrapping_and_symbols():
    name = "A very long plasmid stock identifier with a distinguishing suffix - stock 001"
    rows = [(f"{row}{column}", name, (1 + (column - 1) % 5) * 100)
            for row in "ABCDEFGH" for column in range(1, 13)]
    workbook, _ = workbook_for(rows, 96, stocks={"Stocks": [["Plasmid", "Concentration"], [name, 100]]})
    try:
        for sheet in (workbook["Mix Map (A-D)"], workbook["Mix Map (E-H)"]):
            for cell in grid_positions(sheet).values():
                assert name in cell.value
                explicit_lines = len(cell.value.splitlines())
                assert sheet.row_dimensions[cell.row].height > 14 + 12 * explicit_lines + 10
                assert cell.font.sz == 8
    finally:
        workbook.close()


@pytest.mark.parametrize("plate_type", [48, 96])
@pytest.mark.parametrize("reagent", ["LT1", "L2000"])
def test_printed_maps_preserve_complete_distinct_stock_names(plate_type, reagent):
    names = [
        "Shared backbone construct variant alpha - stock 001",
        "Shared backbone construct variant alpha - stock 002",
        "Shared backbone construct variant beta - stock 001",
        # A real stock may already use the old shortened label as its name.
        "Shared backbone const...",
        "=Shared backbone construct variant alpha - stock 003",
    ]
    last_row = PLATE_FORMATS[plate_type].rows[-1]
    rows = [("A1", names[0], 100), ("A1", names[1], 100),
            ("A2", names[3], 100), (f"{last_row}1", names[2], 100),
            (f"{last_row}2", names[0], 100), (f"{last_row}3", names[4], 100)]
    stocks = {"Stocks": [["Plasmid", "Concentration"]] + [[name, 100] for name in names]}
    workbook, artifact = workbook_for(rows, plate_type, reagent, stocks)
    rich = load_workbook(BytesIO(artifact["bytes"]), rich_text=True)
    try:
        positions = {well: (sheet, cell) for sheet in workbook if sheet.title.startswith("Mix Map")
                     for well, cell in grid_positions(sheet).items()}
        expected = {}
        for well, name, _ in rows:
            expected.setdefault(well, []).append(name)
        for well, stock_names in expected.items():
            sheet, cell = positions[well]
            component_lines = cell.value.splitlines()[1:1 + len(stock_names)]
            assert [line.partition("  ")[2] for line in component_lines] == stock_names
            assert cell.alignment.wrap_text
            assert str(rich[sheet.title][cell.coordinate].value) == cell.value
        assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
        assert set(artifact["details"]["Matched plasmid"]) == set(names)
    finally:
        workbook.close()
        rich.close()


@pytest.mark.parametrize("name", [
    "A long plasmid stock identifier " * 30 + "stock 001",
    "Stock\n" * 40 + "001",
])
def test_maps_reject_stock_names_that_would_exceed_excel_row_height(name):
    stocks = {"Stocks": [["Plasmid", "Concentration"], [name, 100]]}
    with pytest.raises(core.MixMapError) as caught:
        workbook_for([("H12", name, 100)], 96, stocks=stocks)
    assert "too tall" in str(caught.value)
    assert any("plate.csv" in detail and "Mix Map (E-H)" in detail
               and "H12" in detail and "409" in detail
               for detail in caught.value.details)
    assert any("shorter, unique" in fix for fix in caught.value.fixes)


def test_names_with_line_breaks_reserve_height_for_each_displayed_line():
    name = "Stock\n" * 12 + "001"
    stocks = {"Stocks": [["Plasmid", "Concentration"], [name, 100]]}
    workbook, _ = workbook_for([("A1", name, 100)], 96, stocks=stocks)
    try:
        sheet = workbook["Mix Map (A-D)"]
        cell = grid_positions(sheet)["A1"]
        assert name in cell.value
        assert sheet.row_dimensions[cell.row].height >= 12 * len(cell.value.splitlines())
    finally:
        workbook.close()
